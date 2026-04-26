from __future__ import annotations

from copy import deepcopy
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping

from mini_layout_engine.engine.planning_engine import PlanningEngine
from mini_layout_engine.registry.family_registry import FamilyRegistry
from mini_layout_engine.registry.layout_registry import LayoutRegistry
from mini_layout_engine.rendering.fit_utils import (
    emu_to_inches,
    estimate_chars_per_line,
    estimate_line_height_in,
    estimate_multiline_wrapped_lines,
    estimate_text_height_in,
    find_largest_fitting_font,
    find_shrink_to_fit_font,
    measure_text_box,
)
from mini_layout_engine.rendering.image_fit import prepare_image_for_fit_mode
from mini_layout_engine.rendering.pptx_helper import MiniPptxHelper
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Pt


@dataclass
class RenderResult:
    output_path: str
    rendered_slides: int
    skipped_slides: int
    warnings: list[str]

    def to_dict(self) -> dict[str, Any]:
        return {
            "output_path": self.output_path,
            "rendered_slides": self.rendered_slides,
            "skipped_slides": self.skipped_slides,
            "warnings": list(self.warnings),
        }


SECTION_MODE_SPECS: dict[str, dict[str, Any]] = {
    "2x1": {
        "rows": 1,
        "cols": 2,
        "capacity": 2,
        "min_label_pt": 24.0,
        "min_title_pt": 20.0,
        "min_desc_pt": 15.0,
        "max_title_lines": 2,
        "max_desc_lines": 4,
        "allow_desc_truncation": True,
        "max_desc_truncation_ratio": 0.10,
    },
    "2x2": {
        "rows": 2,
        "cols": 2,
        "capacity": 4,
        "min_label_pt": 22.0,
        "min_title_pt": 18.0,
        "min_desc_pt": 13.0,
        "max_title_lines": 2,
        "max_desc_lines": 3,
        "allow_desc_truncation": True,
        "max_desc_truncation_ratio": 0.08,
    },
    "3x2": {
        "rows": 3,
        "cols": 2,
        "capacity": 6,
        "min_label_pt": 20.0,
        "min_title_pt": 16.0,
        "min_desc_pt": 12.0,
        "max_title_lines": 1,
        "max_desc_lines": 2,
        "allow_desc_truncation": False,
        "max_desc_truncation_ratio": 0.0,
    },
}

# Dedicated geometry ratios for single-section full-width rendering.
# This is intentionally separate from template slot-derived ratios so the
# label/title/description composition stays balanced in 1x1 pages.
SECTION_SINGLE_FULLWIDTH_RATIOS: dict[str, tuple[float, float, float, float]] = {
    "label": (0.16, 0.23, 0.10, 0.22),
    "title": (0.27, 0.23, 0.58, 0.22),
    "description": (0.27, 0.53, 0.58, 0.31),
}

SECTION_REASON_CODES = (
    "title_min_font",
    "desc_min_font",
    "title_line_limit",
    "desc_line_limit",
    "desc_truncation_required",
    "label_overflow",
)


@dataclass(frozen=True)
class SectionShapeSlot:
    label_name: str
    title_name: str
    description_name: str
    x_in: float
    y_in: float
    w_in: float
    h_in: float


@dataclass(frozen=True)
class SectionFieldRectRatios:
    label: tuple[float, float, float, float]
    title: tuple[float, float, float, float]
    description: tuple[float, float, float, float]


@dataclass
class SectionFitOutcome:
    passed: bool
    section: dict[str, Any]
    rendered: dict[str, str]
    fonts: dict[str, float]
    lines: dict[str, int]
    reason_codes: list[str]


@dataclass(frozen=True)
class SectionGridGeometry:
    content_left_in: float
    content_top_in: float
    content_width_in: float
    content_height_in: float
    slots: list[SectionShapeSlot]
    field_ratios: SectionFieldRectRatios
    template_fonts_pt: dict[str, float]


class MiniPptxRenderer:
    def __init__(self, templates_dir: str | Path | None = None):
        package_root = Path(__file__).resolve().parents[1]
        self.templates_dir = Path(templates_dir or (package_root / "templates")).resolve()
        self.family_registry = FamilyRegistry.from_specs_dir(
            package_root / "specs" / "families"
        )
        self.layout_registry = LayoutRegistry.from_specs_dir(
            package_root / "specs" / "layouts"
        )

    def render_from_request(
        self,
        request: Mapping[str, Any] | dict[str, Any],
        *,
        output_path: str | Path,
    ) -> RenderResult:
        planning_engine = PlanningEngine.from_default_specs()
        plan_result = planning_engine.plan_deck(request).to_dict()
        return self.render_from_plan(plan_result, output_path=output_path)

    def render_from_plan(
        self,
        plan: Mapping[str, Any] | dict[str, Any],
        *,
        output_path: str | Path,
    ) -> RenderResult:
        payload = dict(plan)
        family_id = str(payload.get("design_family_id", "")).strip()
        if not family_id:
            raise ValueError("Plan payload is missing 'design_family_id'.")

        family = self.family_registry.require(family_id)
        template_name = str(payload.get("template_file") or family.template_file).strip()
        if not template_name:
            raise ValueError(f"Family '{family_id}' has no template_file configured.")

        template_path = self.templates_dir / template_name
        if not template_path.is_file():
            raise FileNotFoundError(f"Template file not found: {template_path}")

        helper = MiniPptxHelper(template_path)
        slides_payload = payload.get("slides", [])
        if not isinstance(slides_payload, list):
            raise ValueError("Plan payload field 'slides' must be a list.")

        warnings: list[str] = []
        prototype_map = self._resolve_prototype_indices(
            family_layout_mapping=family.layout_mapping,
            total_template_slides=helper.slide_count(),
        )

        original_slide_count = helper.slide_count()
        rendered = 0
        skipped = 0

        for slide_entry in slides_payload:
            if not isinstance(slide_entry, Mapping):
                skipped += 1
                warnings.append("Skipped invalid slide entry (not an object).")
                continue

            readiness = str(slide_entry.get("readiness", "")).strip().lower()
            if readiness and readiness != "ready":
                skipped += 1
                warnings.append(
                    f"Skipped slide with layout '{slide_entry.get('layout_id')}' due to readiness='{readiness}'."
                )
                continue

            layout_id = str(slide_entry.get("layout_id", "")).strip()
            if not layout_id:
                skipped += 1
                warnings.append("Skipped slide without layout_id.")
                continue

            prototype_index = prototype_map.get(layout_id)
            if prototype_index is None:
                skipped += 1
                warnings.append(
                    f"Skipped slide with layout '{layout_id}' because family mapping has no prototype index."
                )
                continue

            mapped_layout = slide_entry.get("mapped_layout")
            mapped_layout = mapped_layout if isinstance(mapped_layout, Mapping) else {}
            has_placeholders = isinstance(mapped_layout.get("placeholders"), Mapping)
            content = self._select_content_source(
                slide_entry,
                prefer_input=has_placeholders,
            )

            if self._supports_paginated_table(layout_id, mapped_layout):
                rendered += self._render_paginated_table_slides(
                    helper,
                    prototype_index,
                    content,
                    mapped_layout=mapped_layout,
                    warnings=warnings,
                )
                continue
            if self._supports_adaptive_tabular_data_summary(layout_id, mapped_layout):
                rendered += self._render_adaptive_tabular_data_summary_slides(
                    helper,
                    prototype_index,
                    content,
                    mapped_layout=mapped_layout,
                    warnings=warnings,
                )
                continue
            if self._supports_adaptive_section_grid(layout_id, mapped_layout):
                rendered += self._render_adaptive_section_grid_slides(
                    helper,
                    prototype_index,
                    content,
                    mapped_layout=mapped_layout,
                    warnings=warnings,
                )
                continue
            if self._supports_case_timeline(layout_id, mapped_layout):
                rendered += self._render_case_timeline_slides(
                    helper,
                    prototype_index,
                    content,
                    mapped_layout=mapped_layout,
                    warnings=warnings,
                )
                continue

            helper.duplicate_slide(prototype_index)
            target_index = helper.slide_count() - 1

            self._render_slide_content(
                helper,
                target_index,
                content,
                mapped_layout=mapped_layout,
                layout_id=layout_id,
            )
            if layout_id == "cover_title":
                self._apply_cover_title_shrink_to_fit(
                    helper,
                    target_index,
                    content=content,
                    mapped_layout=mapped_layout,
                )
            if layout_id == "hero_statement":
                self._apply_hero_statement_shrink_to_fit(
                    helper,
                    target_index,
                    content=content,
                    mapped_layout=mapped_layout,
                )
            if layout_id == "title_bullets_with_image":
                self._apply_placeholder_text_bold(
                    helper,
                    target_index,
                    mapped_layout=mapped_layout,
                    placeholder_key="title",
                    is_bold=True,
                )
                self._apply_title_bullets_with_image_title_shrink_to_fit(
                    helper,
                    target_index,
                    content=content,
                    mapped_layout=mapped_layout,
                )
            if layout_id == "two_column_long_text":
                self._apply_mapped_title_shrink_to_fit(
                    helper,
                    target_index,
                    content=content,
                    mapped_layout=mapped_layout,
                    fallback_template_pt=40.0,
                    min_font_pt=12.0,
                )
                self._apply_two_column_long_text_paragraph_grow_to_fit(
                    helper,
                    target_index,
                    mapped_layout=mapped_layout,
                )
            if layout_id == "large_media_title_caption":
                self._apply_mapped_title_shrink_to_fit(
                    helper,
                    target_index,
                    content=content,
                    mapped_layout=mapped_layout,
                    fallback_template_pt=40.0,
                    min_font_pt=12.0,
                )

            image_ref = self._extract_image_ref(content)
            if image_ref:
                image_path = self._resolve_image_path(image_ref)
                if image_path is None:
                    warnings.append(
                        f"Image '{image_ref}' not found for layout '{layout_id}'; kept template image."
                    )
                else:
                    replaced = self._replace_image_for_mapped_placeholder(
                        helper,
                        target_index,
                        mapped_layout=mapped_layout,
                        image_path=image_path,
                    )
                    if not replaced:
                        picture_shapes = helper.sort_shapes_reading_order(
                            helper.picture_shapes(target_index)
                        )
                        if picture_shapes:
                            helper.replace_picture_shape(picture_shapes[0], image_path)
                            replaced = True
                    if not replaced:
                        warnings.append(
                            f"No picture shape available for image '{image_ref}' in layout '{layout_id}'."
                        )

            rendered += 1

        for index in reversed(range(original_slide_count)):
            helper.delete_slide(index)

        helper.assert_all_slides_relationship_integrity()
        helper.assert_all_slides_hyperlink_integrity()
        helper.assert_all_slides_cross_references()
        helper.assert_all_slides_have_no_duplicate_text_regions()
        out_path = Path(output_path).resolve()
        helper.save(out_path)
        return RenderResult(
            output_path=str(out_path),
            rendered_slides=rendered,
            skipped_slides=skipped,
            warnings=warnings,
        )

    def _resolve_prototype_indices(
        self,
        *,
        family_layout_mapping: Mapping[str, Mapping[str, Any]],
        total_template_slides: int,
    ) -> dict[str, int]:
        prototype_map: dict[str, int] = {}
        fallback_index = 0
        for layout_id, mapping in family_layout_mapping.items():
            if not isinstance(mapping, Mapping):
                continue
            explicit = mapping.get("template_slide_index")
            if explicit is not None:
                try:
                    resolved = int(explicit) - 1
                except Exception:
                    resolved = -1
            else:
                resolved = fallback_index
                fallback_index += 1

            if 0 <= resolved < total_template_slides:
                prototype_map[str(layout_id)] = resolved
        return prototype_map

    def _apply_cover_title_shrink_to_fit(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        *,
        content: Mapping[str, Any],
        mapped_layout: Mapping[str, Any],
    ) -> None:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return
        title_mapping = placeholders.get("title")
        if not isinstance(title_mapping, Mapping):
            return

        title_shape_name = str(title_mapping.get("name", "")).strip()
        if not title_shape_name:
            return
        title_shape = helper.get_shape_by_name(slide_index, title_shape_name)
        if title_shape is None or not getattr(title_shape, "has_text_frame", False):
            return

        title_text = str(content.get("title", "") or "").strip()
        if not title_text:
            return

        text_frame = title_shape.text_frame
        margin_left = int(getattr(text_frame, "margin_left", 0) or 0)
        margin_right = int(getattr(text_frame, "margin_right", 0) or 0)
        margin_top = int(getattr(text_frame, "margin_top", 0) or 0)
        margin_bottom = int(getattr(text_frame, "margin_bottom", 0) or 0)
        width_emu = int(title_shape.width) - margin_left - margin_right
        height_emu = int(title_shape.height) - margin_top - margin_bottom
        fit_width_in = max(0.1, emu_to_inches(width_emu))
        fit_height_in = max(0.1, emu_to_inches(height_emu))

        template_font_pt = self._shape_template_font_pt(title_shape, fallback=52.0)
        min_font_pt = 12.0
        fit = find_shrink_to_fit_font(
            title_text,
            width_in=fit_width_in,
            height_in=fit_height_in,
            template_font_pt=float(template_font_pt),
            min_font_pt=min_font_pt,
            step_pt=0.5,
            line_spacing=1.0,
            vertical_padding_in=0.02,
        )
        final_font_pt: float
        if fit is None:
            # Keep cover output stable for extreme titles: never leave template-size overflow.
            final_font_pt = float(min_font_pt)
        else:
            final_font_pt = float(fit.font_size_pt)
        if final_font_pt <= float(template_font_pt):
            self._set_text_shape_font_size(title_shape, final_font_pt)

    def _supports_paginated_table(
        self,
        layout_id: str,
        mapped_layout: Mapping[str, Any],
    ) -> bool:
        if layout_id != "info_table":
            return False
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return False
        rows_mapping = placeholders.get("rows")
        return (
            isinstance(rows_mapping, Mapping)
            and str(rows_mapping.get("type", "")).strip() == "table"
        )

    def _apply_mapped_title_shrink_to_fit(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        *,
        content: Mapping[str, Any],
        mapped_layout: Mapping[str, Any],
        fallback_template_pt: float,
        min_font_pt: float = 12.0,
    ) -> None:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return
        title_mapping = placeholders.get("title")
        if not isinstance(title_mapping, Mapping):
            return

        title_shape_name = str(title_mapping.get("name", "")).strip()
        if not title_shape_name:
            return
        title_shape = helper.get_shape_by_name(slide_index, title_shape_name)
        if title_shape is None or not getattr(title_shape, "has_text_frame", False):
            return

        title_text = str(content.get("title", "") or "").strip()
        if not title_text:
            return

        text_frame = title_shape.text_frame
        margin_left = int(getattr(text_frame, "margin_left", 0) or 0)
        margin_right = int(getattr(text_frame, "margin_right", 0) or 0)
        margin_top = int(getattr(text_frame, "margin_top", 0) or 0)
        margin_bottom = int(getattr(text_frame, "margin_bottom", 0) or 0)
        width_emu = int(title_shape.width) - margin_left - margin_right
        height_emu = int(title_shape.height) - margin_top - margin_bottom
        fit_width_in = max(0.1, emu_to_inches(width_emu))
        fit_height_in = max(0.1, emu_to_inches(height_emu))

        template_font_pt = self._shape_template_font_pt(
            title_shape,
            fallback=float(fallback_template_pt),
        )
        fit = find_shrink_to_fit_font(
            title_text,
            width_in=fit_width_in,
            height_in=fit_height_in,
            template_font_pt=float(template_font_pt),
            min_font_pt=float(min_font_pt),
            step_pt=0.5,
            line_spacing=1.0,
            vertical_padding_in=0.02,
        )
        if fit is None:
            self._set_text_shape_font_size(title_shape, float(min_font_pt))
            return
        if float(fit.font_size_pt) < float(template_font_pt):
            self._set_text_shape_font_size(title_shape, float(fit.font_size_pt))

    def _apply_placeholder_text_bold(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        *,
        mapped_layout: Mapping[str, Any],
        placeholder_key: str,
        is_bold: bool,
    ) -> None:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return
        text_mapping = placeholders.get(str(placeholder_key))
        if not isinstance(text_mapping, Mapping):
            return

        shape_name = str(text_mapping.get("name", "")).strip()
        if not shape_name:
            return
        shape = helper.get_shape_by_name(slide_index, shape_name)
        if shape is None or not getattr(shape, "has_text_frame", False):
            return

        self._set_text_shape_bold(shape, is_bold)

    def _apply_hero_statement_shrink_to_fit(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        *,
        content: Mapping[str, Any],
        mapped_layout: Mapping[str, Any],
    ) -> None:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return
        title_mapping = placeholders.get("title")
        if not isinstance(title_mapping, Mapping):
            return

        title_shape_name = str(title_mapping.get("name", "")).strip()
        if not title_shape_name:
            return
        title_shape = helper.get_shape_by_name(slide_index, title_shape_name)
        if title_shape is None or not getattr(title_shape, "has_text_frame", False):
            return

        title_text = str(content.get("title", "") or "").strip()
        if not title_text:
            return

        text_frame = title_shape.text_frame
        margin_left = int(getattr(text_frame, "margin_left", 0) or 0)
        margin_right = int(getattr(text_frame, "margin_right", 0) or 0)
        margin_top = int(getattr(text_frame, "margin_top", 0) or 0)
        margin_bottom = int(getattr(text_frame, "margin_bottom", 0) or 0)
        width_emu = int(title_shape.width) - margin_left - margin_right
        height_emu = int(title_shape.height) - margin_top - margin_bottom
        fit_width_in = max(0.1, emu_to_inches(width_emu))
        fit_height_in = max(0.1, emu_to_inches(height_emu))

        template_font_pt = self._hero_statement_template_font_pt(title_shape, fallback=52.0)
        min_font_pt = 12.0
        fit = find_shrink_to_fit_font(
            title_text,
            width_in=fit_width_in,
            height_in=fit_height_in,
            template_font_pt=float(template_font_pt),
            min_font_pt=min_font_pt,
            step_pt=0.5,
            line_spacing=1.0,
            vertical_padding_in=0.02,
        )
        if fit is None:
            self._set_text_shape_font_size(title_shape, float(min_font_pt))
            return

        # Preserve template-inherited look when text already fits at baseline.
        normalized = " ".join(title_text.split())
        has_multiple_words = len(normalized.split()) > 1
        wrapped_lines = int(getattr(fit.metrics, "estimated_lines", 1) or 1)
        should_force_explicit_fit = has_multiple_words or wrapped_lines > 1
        if (
            float(fit.font_size_pt) < float(template_font_pt)
            or should_force_explicit_fit
        ):
            self._set_text_shape_font_size(title_shape, float(fit.font_size_pt))

    def _apply_title_bullets_with_image_title_shrink_to_fit(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        *,
        content: Mapping[str, Any],
        mapped_layout: Mapping[str, Any],
    ) -> None:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return
        title_mapping = placeholders.get("title")
        if not isinstance(title_mapping, Mapping):
            return

        title_shape_name = str(title_mapping.get("name", "")).strip()
        if not title_shape_name:
            return
        title_shape = helper.get_shape_by_name(slide_index, title_shape_name)
        if title_shape is None or not getattr(title_shape, "has_text_frame", False):
            return

        title_text = str(content.get("title", "") or "").strip()
        if not title_text:
            return

        text_frame = title_shape.text_frame
        margin_left = int(getattr(text_frame, "margin_left", 0) or 0)
        margin_right = int(getattr(text_frame, "margin_right", 0) or 0)
        margin_top = int(getattr(text_frame, "margin_top", 0) or 0)
        margin_bottom = int(getattr(text_frame, "margin_bottom", 0) or 0)
        width_emu = int(title_shape.width) - margin_left - margin_right
        height_emu = int(title_shape.height) - margin_top - margin_bottom
        fit_width_in = max(0.1, emu_to_inches(width_emu))
        fit_height_in = max(0.1, emu_to_inches(height_emu))

        template_font_pt = self._shape_template_font_pt(title_shape, fallback=44.0)
        min_font_pt = 12.0
        fit = find_shrink_to_fit_font(
            title_text,
            width_in=fit_width_in,
            height_in=fit_height_in,
            template_font_pt=float(template_font_pt),
            min_font_pt=min_font_pt,
            step_pt=0.5,
            line_spacing=1.0,
            vertical_padding_in=0.02,
        )
        if fit is None:
            self._set_text_shape_font_size(title_shape, float(min_font_pt))
            return
        if float(fit.font_size_pt) < float(template_font_pt):
            self._set_text_shape_font_size(title_shape, float(fit.font_size_pt))

    def _hero_statement_template_font_pt(self, shape: Any, *, fallback: float) -> float:
        # Prefer explicit run sizing when present.
        explicit = self._shape_template_font_pt(shape, fallback=fallback)
        if explicit != float(fallback):
            return float(explicit)

        # Hero title in this template can inherit sizing from placeholder style.
        try:
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            def_rpr = shape._element.find(
                ".//a:txBody/a:lstStyle/a:lvl1pPr/a:defRPr",
                namespaces=ns,
            )
            if def_rpr is not None:
                sz = def_rpr.get("sz")
                if sz:
                    return max(1.0, float(sz) / 100.0)
        except Exception:
            pass
        return float(fallback)

    def _apply_two_column_long_text_paragraph_grow_to_fit(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        *,
        mapped_layout: Mapping[str, Any],
    ) -> None:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return
        for field_name in ("left_text", "right_text"):
            mapping = placeholders.get(field_name)
            if not isinstance(mapping, Mapping):
                continue
            shape_name = str(mapping.get("name", "")).strip()
            if not shape_name:
                continue
            shape = helper.get_shape_by_name(slide_index, shape_name)
            if shape is None or not getattr(shape, "has_text_frame", False):
                continue
            self._grow_text_shape_font_size_if_underfilled(shape)

    def _grow_text_shape_font_size_if_underfilled(self, shape: Any) -> None:
        if not getattr(shape, "has_text_frame", False):
            return
        text_frame = shape.text_frame
        text = str(getattr(text_frame, "text", "") or "").strip()
        if not text:
            return

        margin_left = int(getattr(text_frame, "margin_left", 0) or 0)
        margin_right = int(getattr(text_frame, "margin_right", 0) or 0)
        margin_top = int(getattr(text_frame, "margin_top", 0) or 0)
        margin_bottom = int(getattr(text_frame, "margin_bottom", 0) or 0)
        width_emu = int(getattr(shape, "width", 0)) - margin_left - margin_right
        height_emu = int(getattr(shape, "height", 0)) - margin_top - margin_bottom
        width_in = max(0.1, emu_to_inches(width_emu))
        height_in = max(0.1, emu_to_inches(height_emu))

        baseline_font_pt = self._shape_template_font_pt(shape, fallback=14.0)

        max_font_pt = min(
            28.0,
            max(float(baseline_font_pt) + 6.0, float(baseline_font_pt) * 1.45),
        )
        fit = find_largest_fitting_font(
            text,
            width_in=width_in,
            height_in=height_in,
            min_font_pt=float(baseline_font_pt),
            max_font_pt=float(max_font_pt),
            step_pt=0.5,
            line_spacing=1.0,
            vertical_padding_in=0.02,
        )
        if fit is None:
            return
        if float(fit.font_size_pt) > float(baseline_font_pt):
            self._set_text_shape_font_size(shape, float(fit.font_size_pt))

    def _supports_adaptive_tabular_data_summary(
        self,
        layout_id: str,
        mapped_layout: Mapping[str, Any],
    ) -> bool:
        if layout_id != "tabular_data_summary":
            return False
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return False
        table_mapping = placeholders.get("table")
        return (
            isinstance(table_mapping, Mapping)
            and str(table_mapping.get("type", "")).strip() == "table"
        )

    def _render_adaptive_tabular_data_summary_slides(
        self,
        helper: MiniPptxHelper,
        prototype_index: int,
        content: Mapping[str, Any],
        *,
        mapped_layout: Mapping[str, Any],
        warnings: list[str],
    ) -> int:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return 0

        table_mapping = placeholders.get("table")
        if not isinstance(table_mapping, Mapping):
            return 0

        table_rows = [
            dict(row)
            for row in content.get("table", [])
            if isinstance(row, Mapping)
        ]
        header_row, data_rows = self._split_tabular_header_and_data_rows(table_rows)
        active_columns = self._resolve_tabular_active_columns(
            rows=table_rows,
            table_mapping=table_mapping,
        )
        if not active_columns:
            warnings.append(
                "tabular_data_summary has no usable columns in mapping; rendered static table placeholder."
            )
            helper.duplicate_slide(prototype_index)
            target_index = helper.slide_count() - 1
            self._render_slide_content(
                helper,
                target_index,
                content,
                mapped_layout=mapped_layout,
                layout_id="tabular_data_summary",
            )
            self._apply_mapped_title_shrink_to_fit(
                helper,
                target_index,
                content=content,
                mapped_layout=mapped_layout,
                fallback_template_pt=38.0,
                min_font_pt=12.0,
            )
            return 1

        table_mapping_effective = dict(table_mapping)
        table_mapping_effective["columns"] = dict(active_columns)
        # Keep row 0 reserved for repeated header on every continuation page.
        table_mapping_effective["start_row"] = 1
        mapped_layout_effective = dict(mapped_layout)
        placeholders_effective = dict(placeholders)
        placeholders_effective["table"] = table_mapping_effective
        mapped_layout_effective["placeholders"] = placeholders_effective
        table_plan = self._build_table_pagination_plan(
            helper,
            prototype_index,
            table_mapping_effective,
            data_rows,
            warnings=warnings,
        )
        pages = table_plan["pages"]
        if not pages:
            pages = [{"rows": [], "active_row_heights_in": [], "target_table_height_in": 0.0}]

        template_shape = helper.get_shape_by_name(
            prototype_index, str(table_mapping.get("name", ""))
        )
        template_table_height_emu = int(getattr(template_shape, "height", 0) or 0)

        rendered = 0
        for page_index, page in enumerate(pages):
            helper.duplicate_slide(prototype_index)
            target_index = helper.slide_count() - 1
            page_content = dict(content)
            page_content["table"] = page.get("rows", [])
            self._render_slide_content(
                helper,
                target_index,
                page_content,
                mapped_layout=mapped_layout_effective,
                layout_id="tabular_data_summary",
            )
            self._apply_mapped_title_shrink_to_fit(
                helper,
                target_index,
                content=page_content,
                mapped_layout=mapped_layout_effective,
                fallback_template_pt=38.0,
                min_font_pt=12.0,
            )
            self._write_tabular_header_row(
                helper,
                target_index,
                table_mapping_effective,
                header_row=header_row,
                active_columns=active_columns,
            )
            self._apply_tabular_active_columns_geometry(
                helper,
                target_index,
                table_mapping_effective,
                active_columns=active_columns,
            )
            self._apply_table_page_geometry(
                helper,
                target_index,
                table_mapping_effective,
                page["active_row_heights_in"],
                float(table_plan["font_size_pt"]),
                target_table_height_in=float(page.get("target_table_height_in", 0.0)),
            )
            if page_index < len(pages) - 1 and template_table_height_emu > 0:
                page_shape = helper.get_shape_by_name(
                    target_index, str(table_mapping_effective.get("name", ""))
                )
                if page_shape is not None:
                    page_shape.height = int(template_table_height_emu)
                self._expand_tabular_rows_to_shape_height(
                    helper,
                    target_index,
                    table_mapping_effective,
                )
            rendered += 1
        return rendered

    @staticmethod
    def _split_tabular_header_and_data_rows(
        rows: list[dict[str, Any]],
    ) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        if not rows:
            return {}, []
        header_row = dict(rows[0]) if isinstance(rows[0], Mapping) else {}
        data_rows = [dict(row) for row in rows[1:] if isinstance(row, Mapping)]
        return header_row, data_rows

    def _write_tabular_header_row(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        table_mapping: Mapping[str, Any],
        *,
        header_row: Mapping[str, Any],
        active_columns: Mapping[str, int],
    ) -> None:
        shape = helper.get_shape_by_name(slide_index, str(table_mapping.get("name", "")))
        if shape is None or not getattr(shape, "has_table", False):
            return
        table = shape.table
        if len(table.rows) <= 0:
            return

        active_column_indexes = {
            int(column_index)
            for column_index in active_columns.values()
            if 0 <= int(column_index) < len(table.columns)
        }
        for field_name, column_index in active_columns.items():
            idx = int(column_index)
            if not (0 <= idx < len(table.columns)):
                continue
            table.cell(0, idx).text = self._format_placeholder_value(
                header_row.get(str(field_name), "")
            )

        for idx in range(len(table.columns)):
            if idx in active_column_indexes:
                continue
            table.cell(0, idx).text = ""

    def _apply_tabular_active_columns_geometry(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        table_mapping: Mapping[str, Any],
        *,
        active_columns: Mapping[str, int],
    ) -> None:
        shape = helper.get_shape_by_name(slide_index, str(table_mapping.get("name", "")))
        if shape is None or not getattr(shape, "has_table", False):
            return
        table = shape.table
        if len(table.columns) <= 0:
            return

        active_indices = sorted(
            {
                int(column_index)
                for column_index in active_columns.values()
                if 0 <= int(column_index) < len(table.columns)
            }
        )
        if not active_indices:
            return

        keep_count = int(max(active_indices)) + 1
        keep_count = min(max(1, keep_count), len(table.columns))
        # Preserve template visual span by using the original table grid width,
        # not the graphic-frame width.
        total_width = sum(int(column.width) for column in table.columns)
        if total_width <= 0:
            total_width = int(shape.width)
        if total_width <= 0:
            return

        self._prune_table_to_column_count(table, keep_count=keep_count)
        table = shape.table
        if len(table.columns) <= 0:
            return

        current_widths = [int(column.width) for column in table.columns]
        current_total = sum(current_widths)
        if current_total <= 0:
            base = total_width // len(table.columns)
            remainder = total_width - (base * len(table.columns))
            for idx, column in enumerate(table.columns):
                width = base + (1 if idx < remainder else 0)
                column.width = int(max(1, width))
            return

        assigned = 0
        for idx, column in enumerate(table.columns):
            if idx == len(table.columns) - 1:
                width = max(1, total_width - assigned)
            else:
                raw = int(round((float(current_widths[idx]) / float(current_total)) * total_width))
                rows_left_after = len(table.columns) - idx - 1
                max_for_col = max(1, total_width - assigned - rows_left_after)
                width = max(1, min(raw, max_for_col))
                assigned += width
            column.width = int(width)

    def _prune_table_to_column_count(self, table: Any, *, keep_count: int) -> None:
        tbl = getattr(table, "_tbl", None)
        if tbl is None:
            return
        keep = max(1, int(keep_count))
        grid = getattr(tbl, "tblGrid", None)
        if grid is not None:
            grid_cols = list(getattr(grid, "gridCol_lst", []))
            for grid_col in reversed(grid_cols[keep:]):
                grid.remove(grid_col)
        for row in list(getattr(tbl, "tr_lst", [])):
            cells = list(getattr(row, "tc_lst", []))
            for cell in reversed(cells[keep:]):
                row.remove(cell)

    def _expand_tabular_rows_to_shape_height(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        table_mapping: Mapping[str, Any],
    ) -> None:
        shape = helper.get_shape_by_name(slide_index, str(table_mapping.get("name", "")))
        if shape is None or not getattr(shape, "has_table", False):
            return
        table = shape.table
        if len(table.rows) <= 0:
            return

        start_row = int(table_mapping.get("start_row", 0) or 0)
        start_row = max(0, min(start_row, len(table.rows)))
        if start_row >= len(table.rows):
            return

        fixed_height_emu = sum(int(table.rows[idx].height) for idx in range(start_row))
        target_body_emu = int(shape.height) - fixed_height_emu
        active_rows = [table.rows[idx] for idx in range(start_row, len(table.rows))]
        if not active_rows or target_body_emu <= 0:
            return

        current_body_emu = sum(int(row.height) for row in active_rows)
        if current_body_emu <= 0:
            even = max(1, target_body_emu // len(active_rows))
            remaining = int(target_body_emu)
            for idx, row in enumerate(active_rows):
                if idx == len(active_rows) - 1:
                    row.height = int(max(1, remaining))
                else:
                    row.height = int(max(1, even))
                    remaining -= int(row.height)
            return

        scale = float(target_body_emu) / float(current_body_emu)
        assigned = 0
        for idx, row in enumerate(active_rows):
            if idx == len(active_rows) - 1:
                row.height = int(max(1, target_body_emu - assigned))
                continue
            rows_left_after = len(active_rows) - idx - 1
            raw = int(round(float(row.height) * scale))
            max_for_row = max(1, target_body_emu - assigned - rows_left_after)
            height_emu = max(1, min(raw, max_for_row))
            row.height = int(height_emu)
            assigned += int(height_emu)

    def _resolve_tabular_active_columns(
        self,
        *,
        rows: list[dict[str, Any]],
        table_mapping: Mapping[str, Any],
    ) -> dict[str, int]:
        columns_raw = table_mapping.get("columns")
        if not isinstance(columns_raw, Mapping):
            return {}

        ordered: list[tuple[str, int]] = []
        for field_name, column_index in columns_raw.items():
            try:
                idx = int(column_index)
            except Exception:
                continue
            if idx < 0:
                continue
            ordered.append((str(field_name), idx))
        if not ordered:
            return {}
        ordered.sort(key=lambda item: item[1])

        last_populated_position = -1
        for position, (field_name, _) in enumerate(ordered):
            if self._tabular_column_has_content(rows, field_name):
                last_populated_position = position

        if last_populated_position >= 0:
            target_count = last_populated_position + 1
        else:
            target_count = 0

        minimum_count = 2 if len(ordered) >= 2 else len(ordered)
        target_count = max(minimum_count, target_count)
        target_count = min(4, target_count, len(ordered))
        if target_count <= 0:
            return {}

        return {field: idx for field, idx in ordered[:target_count]}

    def _tabular_column_has_content(
        self,
        rows: list[dict[str, Any]],
        field_name: str,
    ) -> bool:
        for row in rows:
            if self._tabular_value_has_content(row.get(field_name)):
                return True
        return False

    @staticmethod
    def _tabular_value_has_content(value: Any) -> bool:
        if value is None:
            return False
        if isinstance(value, str):
            return bool(value.strip())
        if isinstance(value, (list, tuple)):
            return any(MiniPptxRenderer._tabular_value_has_content(item) for item in value)
        return True

    def _supports_adaptive_section_grid(
        self,
        layout_id: str,
        mapped_layout: Mapping[str, Any],
    ) -> bool:
        if layout_id != "section_grid":
            return False
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return False
        sections_mapping = placeholders.get("sections")
        if not isinstance(sections_mapping, Mapping):
            return False
        return str(sections_mapping.get("type", "")).strip() == "repeated_group"

    def _supports_case_timeline(
        self,
        layout_id: str,
        mapped_layout: Mapping[str, Any],
    ) -> bool:
        if layout_id != "case_timeline":
            return False
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return False
        steps_mapping = placeholders.get("steps")
        if not isinstance(steps_mapping, Mapping):
            return False
        return str(steps_mapping.get("type", "")).strip() == "timeline_steps_dynamic"

    def _render_case_timeline_slides(
        self,
        helper: MiniPptxHelper,
        prototype_index: int,
        content: Mapping[str, Any],
        *,
        mapped_layout: Mapping[str, Any],
        warnings: list[str],
    ) -> int:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return 0
        title_mapping = placeholders.get("title")
        title_shape_name = ""
        if isinstance(title_mapping, Mapping):
            title_shape_name = str(title_mapping.get("name", "")).strip()
        steps_mapping = placeholders.get("steps")
        if not isinstance(steps_mapping, Mapping):
            return 0

        prototypes_raw = steps_mapping.get("item_prototypes")
        if not isinstance(prototypes_raw, list) or not prototypes_raw:
            helper.duplicate_slide(prototype_index)
            target_index = helper.slide_count() - 1
            self._render_slide_content(
                helper,
                target_index,
                content,
                mapped_layout=mapped_layout,
                layout_id="case_timeline",
            )
            self._apply_mapped_title_shrink_to_fit(
                helper,
                target_index,
                content=content,
                mapped_layout=mapped_layout,
                fallback_template_pt=40.0,
                min_font_pt=12.0,
            )
            warnings.append("case_timeline missing item_prototypes; rendered static slide.")
            return 1

        slots: list[dict[str, str]] = []
        for entry in prototypes_raw:
            if not isinstance(entry, Mapping):
                continue
            marker = str(entry.get("marker", "")).strip()
            title = str(entry.get("title", "")).strip()
            description = str(entry.get("description", "")).strip()
            if not marker or not title or not description:
                continue
            slot_position = str(entry.get("position", "")).strip().lower()
            if slot_position not in {"top", "bottom"}:
                slot_position = ""
            slots.append(
                {
                    "marker": marker,
                    "title": title,
                    "description": description,
                    "position": slot_position,
                }
            )
        if not slots:
            helper.duplicate_slide(prototype_index)
            target_index = helper.slide_count() - 1
            self._render_slide_content(
                helper,
                target_index,
                content,
                mapped_layout=mapped_layout,
                layout_id="case_timeline",
            )
            self._apply_mapped_title_shrink_to_fit(
                helper,
                target_index,
                content=content,
                mapped_layout=mapped_layout,
                fallback_template_pt=40.0,
                min_font_pt=12.0,
            )
            warnings.append("case_timeline has invalid item_prototypes; rendered static slide.")
            return 1

        axis_mapping = steps_mapping.get("axis")
        axis_mapping = axis_mapping if isinstance(axis_mapping, Mapping) else {}
        position_pattern_raw = axis_mapping.get("position_pattern")
        position_pattern: list[str] = []
        if isinstance(position_pattern_raw, list):
            for value in position_pattern_raw:
                normalized = str(value or "").strip().lower()
                if normalized in {"top", "bottom"}:
                    position_pattern.append(normalized)
        if not position_pattern:
            position_pattern = ["top", "bottom"]

        line_segment_names_raw = axis_mapping.get("line_segments")
        line_segment_names = [
            str(name).strip()
            for name in line_segment_names_raw
            if str(name).strip()
        ] if isinstance(line_segment_names_raw, list) else []

        steps = self._coerce_timeline_steps(content.get("steps"))
        chunk_sizes = self._split_timeline_chunk_sizes(
            total_steps=len(steps),
            max_per_slide=len(slots),
        )
        if not chunk_sizes:
            chunk_sizes = [0]

        rendered = 0
        cursor = 0
        title_text = str(content.get("title", "") or "")
        for page_index, chunk_size in enumerate(chunk_sizes):
            page_steps_raw = steps[cursor: cursor + chunk_size]
            page_steps: list[dict[str, Any]] = []
            for local_index, step in enumerate(page_steps_raw):
                global_index = cursor + local_index + 1
                step_index = str(step.get("index", "") or "").strip()
                resolved = dict(step)
                resolved["index"] = step_index or f"{global_index:02d}"
                page_steps.append(resolved)

            helper.duplicate_slide(prototype_index)
            target_index = helper.slide_count() - 1
            if title_shape_name:
                title_shape = helper.get_shape_by_name(target_index, title_shape_name)
                if title_shape is not None:
                    helper.replace_text_preserve_format(title_shape, title_text)
            self._apply_mapped_title_shrink_to_fit(
                helper,
                target_index,
                content=content,
                mapped_layout=mapped_layout,
                fallback_template_pt=40.0,
                min_font_pt=12.0,
            )

            assignments = self._assign_timeline_steps_to_slots(
                steps=page_steps,
                slots=slots,
                position_pattern=position_pattern,
            )
            self._render_case_timeline_page(
                helper,
                slide_index=target_index,
                slots=slots,
                assignments=assignments,
                line_segment_names=line_segment_names,
                has_next=page_index < len(chunk_sizes) - 1,
                continuation_page=page_index > 0,
            )
            rendered += 1
            cursor += chunk_size
        return rendered

    @staticmethod
    def _coerce_timeline_steps(value: Any) -> list[dict[str, Any]]:
        if not isinstance(value, list):
            return []
        items: list[dict[str, Any]] = []
        for entry in value:
            if not isinstance(entry, Mapping):
                continue
            position = str(entry.get("position", "") or "").strip().lower()
            if position not in {"top", "bottom"}:
                position = ""
            items.append(
                {
                    "index": str(entry.get("index", "") or "").strip(),
                    "title": str(entry.get("title", "") or "").strip(),
                    "description": str(entry.get("description", "") or "").strip(),
                    "position": position,
                }
            )
        return items

    @staticmethod
    def _split_timeline_chunk_sizes(*, total_steps: int, max_per_slide: int) -> list[int]:
        total = max(0, int(total_steps))
        max_items = max(1, int(max_per_slide))
        if total == 0:
            return [0]

        chunks: list[int] = []
        remaining = total
        while remaining > 0:
            if remaining <= max_items:
                chunks.append(remaining)
                break
            take = max_items
            if remaining - take == 1:
                take = max(1, max_items - 1)
            chunks.append(take)
            remaining -= take
        return chunks

    def _assign_timeline_steps_to_slots(
        self,
        *,
        steps: list[dict[str, Any]],
        slots: list[dict[str, str]],
        position_pattern: list[str],
    ) -> list[dict[str, Any]]:
        assignments: list[dict[str, Any]] = []
        next_slot = 0
        for local_index, step in enumerate(steps):
            default_position = position_pattern[local_index % len(position_pattern)]
            desired = str(step.get("position", "") or "").strip().lower() or default_position
            if desired not in {"top", "bottom"}:
                desired = default_position

            slot_index = self._pick_timeline_slot_index(
                slots=slots,
                start_index=next_slot,
                desired_position=desired,
            )
            if slot_index is None:
                break
            assignments.append({"slot_index": slot_index, "step": step})
            next_slot = slot_index + 1
        return assignments

    @staticmethod
    def _pick_timeline_slot_index(
        *,
        slots: list[dict[str, str]],
        start_index: int,
        desired_position: str,
    ) -> int | None:
        for idx in range(start_index, len(slots)):
            slot_position = str(slots[idx].get("position", "")).strip().lower()
            if slot_position == desired_position:
                return idx
        for idx in range(start_index, len(slots)):
            return idx
        return None

    def _render_case_timeline_page(
        self,
        helper: MiniPptxHelper,
        *,
        slide_index: int,
        slots: list[dict[str, str]],
        assignments: list[dict[str, Any]],
        line_segment_names: list[str],
        has_next: bool,
        continuation_page: bool,
    ) -> None:
        assignment_by_slot = {
            int(entry.get("slot_index")): entry.get("step", {})
            for entry in assignments
            if isinstance(entry, Mapping) and entry.get("slot_index") is not None
        }
        visible_slots = sorted(idx for idx in assignment_by_slot.keys() if 0 <= int(idx) < len(slots))
        visible_step_count = len(visible_slots)

        for slot_index, slot in enumerate(slots):
            step = assignment_by_slot.get(slot_index)
            marker_shape = helper.get_shape_by_name(slide_index, str(slot.get("marker", "")))
            title_shape = helper.get_shape_by_name(slide_index, str(slot.get("title", "")))
            desc_shape = helper.get_shape_by_name(slide_index, str(slot.get("description", "")))
            slot_position = str(slot.get("position", "")).strip().lower()
            if not isinstance(step, Mapping):
                for shape in (marker_shape, title_shape, desc_shape):
                    if shape is not None:
                        helper.remove_shape(shape)
                continue

            if marker_shape is not None:
                helper.replace_text_preserve_format(marker_shape, str(step.get("index", "") or ""))
            if title_shape is not None:
                helper.replace_text_preserve_format(title_shape, str(step.get("title", "") or ""))
                self._apply_case_timeline_bottom_title_anchor(
                    title_shape=title_shape,
                    slot_position=slot_position,
                )
                self._apply_case_timeline_text_fit(
                    text_shape=title_shape,
                    min_font_pt=8.0,
                    hard_min_font_pt=5.5,
                    fallback_template_pt=24.0,
                )
            if desc_shape is not None:
                helper.replace_text_preserve_format(desc_shape, str(step.get("description", "") or ""))
                self._apply_case_timeline_text_fit(
                    text_shape=desc_shape,
                    min_font_pt=7.0,
                    hard_min_font_pt=5.5,
                    fallback_template_pt=14.0,
                )

        self._apply_case_timeline_line_segments(
            helper,
            slide_index=slide_index,
            line_segment_names=line_segment_names,
            visible_step_count=visible_step_count,
            has_next=has_next,
            continuation_page=continuation_page,
        )
        self._ensure_case_timeline_markers_above_line(
            helper,
            slide_index=slide_index,
            slots=slots,
            visible_slots=visible_slots,
        )

    def _ensure_case_timeline_markers_above_line(
        self,
        helper: MiniPptxHelper,
        *,
        slide_index: int,
        slots: list[dict[str, str]],
        visible_slots: list[int],
    ) -> None:
        if not visible_slots:
            return
        slide = helper.get_slide(slide_index)
        sp_tree = slide.shapes._spTree
        for slot_index in visible_slots:
            if not (0 <= int(slot_index) < len(slots)):
                continue
            marker_name = str(slots[int(slot_index)].get("marker", "")).strip()
            if not marker_name:
                continue
            marker_shape = helper.get_shape_by_name(slide_index, marker_name)
            if marker_shape is None:
                continue
            element = marker_shape._element
            parent = element.getparent()
            if parent is None:
                continue
            parent.remove(element)
            sp_tree.insert_element_before(element, "p:extLst")

    def _apply_case_timeline_bottom_title_anchor(
        self,
        *,
        title_shape: Any,
        slot_position: str,
    ) -> None:
        if str(slot_position).strip().lower() != "bottom":
            return
        if not getattr(title_shape, "has_text_frame", False):
            return
        text_frame = title_shape.text_frame
        text_frame.word_wrap = True
        text_frame.vertical_anchor = MSO_ANCHOR.TOP

    def _apply_case_timeline_text_fit(
        self,
        *,
        text_shape: Any,
        min_font_pt: float,
        hard_min_font_pt: float,
        fallback_template_pt: float,
    ) -> None:
        if not getattr(text_shape, "has_text_frame", False):
            return
        text_frame = text_shape.text_frame
        raw_text = str(getattr(text_frame, "text", "") or "").strip()
        if not raw_text:
            return
        text_frame.word_wrap = True

        margin_left = int(getattr(text_frame, "margin_left", 0) or 0)
        margin_right = int(getattr(text_frame, "margin_right", 0) or 0)
        margin_top = int(getattr(text_frame, "margin_top", 0) or 0)
        margin_bottom = int(getattr(text_frame, "margin_bottom", 0) or 0)
        width_emu = int(getattr(text_shape, "width", 0)) - margin_left - margin_right
        height_emu = int(getattr(text_shape, "height", 0)) - margin_top - margin_bottom
        fit_width_in = max(0.1, emu_to_inches(width_emu))
        fit_height_in = max(0.1, emu_to_inches(height_emu))

        template_font_pt = self._shape_template_font_pt(
            text_shape,
            fallback=float(fallback_template_pt),
        )
        fit = find_shrink_to_fit_font(
            raw_text,
            width_in=fit_width_in,
            height_in=fit_height_in,
            template_font_pt=float(template_font_pt),
            min_font_pt=float(min_font_pt),
            step_pt=0.5,
            line_spacing=1.0,
            vertical_padding_in=0.02,
        )
        if fit is not None:
            chosen_pt = self._refine_case_timeline_font_size_for_safety(
                text=raw_text,
                width_in=fit_width_in,
                height_in=fit_height_in,
                start_font_pt=float(fit.font_size_pt),
                hard_min_font_pt=float(hard_min_font_pt),
            )
            if float(chosen_pt) <= float(template_font_pt):
                self._set_text_shape_font_size(text_shape, float(chosen_pt))
            return

        # Extreme-content fallback: prioritize in-shape validity over readability.
        fallback_fit = find_shrink_to_fit_font(
            raw_text,
            width_in=fit_width_in,
            height_in=fit_height_in,
            template_font_pt=float(min_font_pt),
            min_font_pt=float(hard_min_font_pt),
            step_pt=0.5,
            line_spacing=1.0,
            vertical_padding_in=0.02,
        )
        if fallback_fit is not None:
            chosen_pt = self._refine_case_timeline_font_size_for_safety(
                text=raw_text,
                width_in=fit_width_in,
                height_in=fit_height_in,
                start_font_pt=float(fallback_fit.font_size_pt),
                hard_min_font_pt=float(hard_min_font_pt),
            )
            self._set_text_shape_font_size(text_shape, float(chosen_pt))
            return

        self._set_text_shape_font_size(text_shape, float(hard_min_font_pt))

    def _refine_case_timeline_font_size_for_safety(
        self,
        *,
        text: str,
        width_in: float,
        height_in: float,
        start_font_pt: float,
        hard_min_font_pt: float,
    ) -> float:
        candidate = max(float(hard_min_font_pt), float(start_font_pt))
        while candidate >= float(hard_min_font_pt) - 1e-6:
            metrics = measure_text_box(
                text,
                width_in=width_in,
                height_in=height_in,
                font_size_pt=candidate,
                line_spacing=1.0,
                vertical_padding_in=0.03,
                average_char_width_factor=0.58,
            )
            if bool(metrics.fits):
                return float(candidate)
            candidate -= 0.5
        return float(hard_min_font_pt)

    def _apply_case_timeline_line_segments(
        self,
        helper: MiniPptxHelper,
        *,
        slide_index: int,
        line_segment_names: list[str],
        visible_step_count: int,
        has_next: bool,
        continuation_page: bool,
    ) -> None:
        if not line_segment_names:
            return
        segment_shapes = []
        for name in line_segment_names:
            shape = helper.get_shape_by_name(slide_index, str(name))
            if shape is not None:
                segment_shapes.append(shape)
        if not segment_shapes:
            return

        count = max(0, int(visible_step_count))
        if count <= 0:
            for shape in segment_shapes:
                helper.remove_shape(shape)
            return

        if continuation_page and segment_shapes:
            first_segment = segment_shapes[0]
            original_left = int(getattr(first_segment, "left", 0))
            if original_left > 0:
                first_segment.width = int(getattr(first_segment, "width", 0)) + original_left
                first_segment.left = 0

        if has_next:
            extension_idx = min(len(segment_shapes) - 1, max(0, count - 1))
            for idx, shape in enumerate(segment_shapes):
                if idx > extension_idx:
                    helper.remove_shape(shape)
                    continue
                if idx == extension_idx:
                    slide_right = int(helper.prs.slide_width)
                    shape.width = max(0, slide_right - int(shape.left))
            return

        keep_count = min(len(segment_shapes), max(0, count - 1))
        for idx, shape in enumerate(segment_shapes):
            if idx >= keep_count:
                helper.remove_shape(shape)

    def _render_adaptive_section_grid_slides(
        self,
        helper: MiniPptxHelper,
        prototype_index: int,
        content: Mapping[str, Any],
        *,
        mapped_layout: Mapping[str, Any],
        warnings: list[str],
    ) -> int:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return 0
        sections_mapping = placeholders.get("sections")
        if not isinstance(sections_mapping, Mapping):
            return 0
        title_mapping = placeholders.get("title")
        title_shape_name = ""
        if isinstance(title_mapping, Mapping):
            title_shape_name = str(title_mapping.get("name", "")).strip()

        sections = self._coerce_section_items(content.get("sections", []))
        geometry = self._extract_section_grid_geometry(
            helper,
            prototype_index=prototype_index,
            sections_mapping=sections_mapping,
            warnings=warnings,
        )
        if geometry is None:
            helper.duplicate_slide(prototype_index)
            target_index = helper.slide_count() - 1
            self._render_slide_content(
                helper,
                target_index,
                content,
                mapped_layout=mapped_layout,
                layout_id="section_grid",
            )
            self._apply_mapped_title_shrink_to_fit(
                helper,
                target_index,
                content=content,
                mapped_layout=mapped_layout,
                fallback_template_pt=40.0,
                min_font_pt=12.0,
            )
            warnings.append(
                "section_grid adaptive mode unavailable; rendered with static repeated-group mapping."
            )
            return 1

        plan = self._plan_section_grid_pages(
            sections=sections,
            geometry=geometry,
            warnings=warnings,
        )
        diagnostics = plan.get("diagnostics", [])
        if isinstance(diagnostics, list):
            for diag in diagnostics:
                if isinstance(diag, Mapping):
                    warnings.append("section_grid_diagnostic: " + json.dumps(dict(diag), ensure_ascii=False))
        pages = plan.get("pages", [])
        if not isinstance(pages, list) or not pages:
            pages = [{"mode": "2x2", "sections": []}]

        rendered = 0
        for page in pages:
            mode = str(page.get("mode", "2x2"))
            page_sections = page.get("sections", [])
            if not isinstance(page_sections, list):
                page_sections = []

            helper.duplicate_slide(prototype_index)
            target_index = helper.slide_count() - 1
            self._render_section_grid_page(
                helper,
                slide_index=target_index,
                title_shape_name=title_shape_name,
                title_text=str(content.get("title", "") or ""),
                sections_mapping=sections_mapping,
                geometry=geometry,
                mode=mode,
                page_sections=page_sections,
                forced_page_size=page.get("forced_page_size"),
            )
            self._apply_mapped_title_shrink_to_fit(
                helper,
                target_index,
                content=content,
                mapped_layout=mapped_layout,
                fallback_template_pt=40.0,
                min_font_pt=12.0,
            )
            rendered += 1
        return rendered

    def _coerce_section_items(self, value: Any) -> list[dict[str, Any]]:
        if not isinstance(value, list):
            return []
        items: list[dict[str, Any]] = []
        for item in value:
            if not isinstance(item, Mapping):
                continue
            items.append(
                {
                    "label": str(item.get("label", "") or "").strip(),
                    "title": str(item.get("title", "") or "").strip(),
                    "description": str(item.get("description", "") or "").strip(),
                }
            )
        return items

    def _extract_section_grid_geometry(
        self,
        helper: MiniPptxHelper,
        *,
        prototype_index: int,
        sections_mapping: Mapping[str, Any],
        warnings: list[str],
    ) -> SectionGridGeometry | None:
        items = sections_mapping.get("items")
        if not isinstance(items, list) or len(items) < 4:
            warnings.append("section_grid placeholders.items must contain at least 4 blocks.")
            return None

        slot_entries: list[SectionShapeSlot] = []
        for entry in items[:4]:
            if not isinstance(entry, Mapping):
                warnings.append("section_grid item mapping is invalid.")
                return None
            label_name = str(entry.get("label", "")).strip()
            title_name = str(entry.get("title", "")).strip()
            description_name = str(entry.get("description", "")).strip()
            label_shape = helper.get_shape_by_name(prototype_index, label_name)
            title_shape = helper.get_shape_by_name(prototype_index, title_name)
            desc_shape = helper.get_shape_by_name(prototype_index, description_name)
            if label_shape is None or title_shape is None or desc_shape is None:
                warnings.append(
                    "section_grid adaptive mode could not find all mapped shapes on prototype slide."
                )
                return None
            left = min(label_shape.left, title_shape.left, desc_shape.left)
            top = min(label_shape.top, title_shape.top, desc_shape.top)
            right = max(
                label_shape.left + label_shape.width,
                title_shape.left + title_shape.width,
                desc_shape.left + desc_shape.width,
            )
            bottom = max(
                label_shape.top + label_shape.height,
                title_shape.top + title_shape.height,
                desc_shape.top + desc_shape.height,
            )
            slot_entries.append(
                SectionShapeSlot(
                    label_name=label_name,
                    title_name=title_name,
                    description_name=description_name,
                    x_in=emu_to_inches(left),
                    y_in=emu_to_inches(top),
                    w_in=emu_to_inches(right - left),
                    h_in=emu_to_inches(bottom - top),
                )
            )

        content_left = min(slot.x_in for slot in slot_entries)
        content_top = min(slot.y_in for slot in slot_entries)
        content_right = max(slot.x_in + slot.w_in for slot in slot_entries)
        content_bottom = max(slot.y_in + slot.h_in for slot in slot_entries)

        ref = slot_entries[0]
        label_ref = helper.get_shape_by_name(prototype_index, ref.label_name)
        title_ref = helper.get_shape_by_name(prototype_index, ref.title_name)
        desc_ref = helper.get_shape_by_name(prototype_index, ref.description_name)
        if label_ref is None or title_ref is None or desc_ref is None:
            return None

        ref_left_emu = int(round(ref.x_in * 914400))
        ref_top_emu = int(round(ref.y_in * 914400))
        ref_w_emu = int(round(ref.w_in * 914400))
        ref_h_emu = int(round(ref.h_in * 914400))
        field_ratios = SectionFieldRectRatios(
            label=self._shape_rect_ratios(label_ref, ref_left_emu, ref_top_emu, ref_w_emu, ref_h_emu),
            title=self._shape_rect_ratios(title_ref, ref_left_emu, ref_top_emu, ref_w_emu, ref_h_emu),
            description=self._shape_rect_ratios(desc_ref, ref_left_emu, ref_top_emu, ref_w_emu, ref_h_emu),
        )

        template_fonts = {
            "label": self._shape_template_font_pt(label_ref, fallback=26.0),
            "title": self._shape_template_font_pt(title_ref, fallback=22.0),
            "description": self._shape_template_font_pt(desc_ref, fallback=16.0),
        }
        return SectionGridGeometry(
            content_left_in=content_left,
            content_top_in=content_top,
            content_width_in=max(0.1, content_right - content_left),
            content_height_in=max(0.1, content_bottom - content_top),
            slots=slot_entries,
            field_ratios=field_ratios,
            template_fonts_pt=template_fonts,
        )

    def _shape_rect_ratios(
        self,
        shape: Any,
        parent_left_emu: int,
        parent_top_emu: int,
        parent_w_emu: int,
        parent_h_emu: int,
    ) -> tuple[float, float, float, float]:
        safe_w = max(1, int(parent_w_emu))
        safe_h = max(1, int(parent_h_emu))
        x = (int(shape.left) - parent_left_emu) / safe_w
        y = (int(shape.top) - parent_top_emu) / safe_h
        w = int(shape.width) / safe_w
        h = int(shape.height) / safe_h
        return (
            max(0.0, float(x)),
            max(0.0, float(y)),
            max(0.01, float(w)),
            max(0.01, float(h)),
        )

    def _shape_template_font_pt(self, shape: Any, *, fallback: float) -> float:
        sizes: list[float] = []
        if not getattr(shape, "has_text_frame", False):
            return float(fallback)
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                size = getattr(run.font, "size", None)
                if size is not None:
                    sizes.append(float(size.pt))
        if not sizes:
            return float(fallback)
        return min(sizes)

    def _plan_section_grid_pages(
        self,
        *,
        sections: list[dict[str, Any]],
        geometry: SectionGridGeometry,
        warnings: list[str],
    ) -> dict[str, Any]:
        n = len(sections)
        diagnostics: list[dict[str, Any]] = []

        if n <= 2:
            attempt = self._evaluate_section_mode("2x1", sections, geometry)
            if attempt["ok"]:
                return {"pages": [{"mode": "2x1", "sections": attempt["sections"]}], "diagnostics": diagnostics}
            if n == 2:
                diagnostics.append(
                    self._build_mode_rejection_diag(
                        rejected_mode="2x1",
                        attempt=attempt,
                        fallback_selected="2x1_continuation",
                    )
                )
                warnings.append("section_grid: 2x1 rejected for n=2; using 2x1 continuation.")
                pages = self._paginate_sections_by_mode(
                    sections=sections,
                    mode="2x1",
                    geometry=geometry,
                    forced_page_size=1,
                    diagnostics=diagnostics,
                )
                return {"pages": pages, "diagnostics": diagnostics}
            diagnostics.append(
                self._build_mode_rejection_diag(
                    rejected_mode="2x1",
                    attempt=attempt,
                    fallback_selected="2x1_single",
                )
            )
            warnings.append("section_grid: 2x1 thresholds not fully met for single section; rendered with bounded fallback.")
            return {"pages": [{"mode": "2x1", "sections": attempt["sections"]}], "diagnostics": diagnostics}

        if n <= 4:
            attempt = self._evaluate_section_mode("2x2", sections, geometry)
            if attempt["ok"]:
                return {"pages": [{"mode": "2x2", "sections": attempt["sections"]}], "diagnostics": diagnostics}
            diagnostics.append(
                self._build_mode_rejection_diag(
                    rejected_mode="2x2",
                    attempt=attempt,
                    fallback_selected="2x2_continuation",
                )
            )
            warnings.append("section_grid: 2x2 fit rejected; using 2x2 continuation fallback.")
            pages = self._paginate_sections_by_mode(
                sections=sections,
                mode="2x2",
                geometry=geometry,
                diagnostics=diagnostics,
            )
            return {"pages": pages, "diagnostics": diagnostics}

        if n <= 6:
            attempt = self._evaluate_section_mode("3x2", sections, geometry)
            if attempt["ok"]:
                return {"pages": [{"mode": "3x2", "sections": attempt["sections"]}], "diagnostics": diagnostics}
            diagnostics.append(
                self._build_mode_rejection_diag(
                    rejected_mode="3x2",
                    attempt=attempt,
                    fallback_selected="2x2_continuation",
                )
            )
            warnings.append("section_grid: 3x2 rejected; using 2x2 continuation fallback.")
            pages = self._paginate_sections_by_mode(
                sections=sections,
                mode="2x2",
                geometry=geometry,
                diagnostics=diagnostics,
            )
            return {"pages": pages, "diagnostics": diagnostics}

        warnings.append("section_grid: section count > 6; using deterministic 2x2 continuation.")
        pages = self._paginate_sections_by_mode(
            sections=sections,
            mode="2x2",
            geometry=geometry,
            diagnostics=diagnostics,
        )
        return {"pages": pages, "diagnostics": diagnostics}

    def _build_mode_rejection_diag(
        self,
        *,
        rejected_mode: str,
        attempt: Mapping[str, Any],
        fallback_selected: str,
    ) -> dict[str, Any]:
        failed_indices = []
        reason_codes: list[str] = []
        outcomes = attempt.get("outcomes", [])
        if isinstance(outcomes, list):
            for idx, outcome in enumerate(outcomes):
                if isinstance(outcome, SectionFitOutcome):
                    if not outcome.passed:
                        failed_indices.append(idx)
                        reason_codes.extend(outcome.reason_codes)
        unique_reasons = sorted(
            reason
            for reason in set(reason_codes)
            if reason in SECTION_REASON_CODES
        )
        return {
            "layout_id": "section_grid",
            "rejected_mode": rejected_mode,
            "reason_codes": unique_reasons,
            "failed_section_indices": failed_indices,
            "fallback_selected": fallback_selected,
        }

    def _paginate_sections_by_mode(
        self,
        *,
        sections: list[dict[str, Any]],
        mode: str,
        geometry: SectionGridGeometry,
        diagnostics: list[dict[str, Any]],
        forced_page_size: int | None = None,
    ) -> list[dict[str, Any]]:
        mode_spec = SECTION_MODE_SPECS[mode]
        max_per_page = int(forced_page_size or mode_spec["capacity"])
        pages: list[dict[str, Any]] = []
        cursor = 0
        while cursor < len(sections):
            chunk_raw = sections[cursor: cursor + max_per_page]
            attempt = self._evaluate_section_mode(mode, chunk_raw, geometry)
            if not attempt["ok"] and max_per_page > 1:
                max_per_page -= 1
                diagnostics.append(
                    self._build_mode_rejection_diag(
                        rejected_mode=mode,
                        attempt=attempt,
                        fallback_selected=f"{mode}_continuation_smaller_chunk_{max_per_page}",
                    )
                )
                continue
            pages.append(
                {
                    "mode": mode,
                    "forced_page_size": int(forced_page_size) if forced_page_size is not None else None,
                    "sections": attempt["sections"] if isinstance(attempt.get("sections"), list) else chunk_raw,
                }
            )
            cursor += len(chunk_raw)
        if not pages:
            pages = [{"mode": mode, "sections": []}]
        return pages

    def _evaluate_section_mode(
        self,
        mode: str,
        sections: list[dict[str, Any]],
        geometry: SectionGridGeometry,
    ) -> dict[str, Any]:
        mode_spec = SECTION_MODE_SPECS[mode]
        rows = int(mode_spec["rows"])
        cols = int(mode_spec["cols"])
        outcomes: list[SectionFitOutcome] = []
        rendered_sections: list[dict[str, Any]] = []
        ok = True
        for idx, section in enumerate(sections):
            field_rects = self._section_mode_field_rects(
                mode=mode,
                geometry=geometry,
                section_index=idx,
            )
            outcome = self._evaluate_single_section_fit(
                section=section,
                field_rects=field_rects,
                mode_spec=mode_spec,
                geometry=geometry,
            )
            outcomes.append(outcome)
            rendered_sections.append(
                {
                    "label": outcome.rendered.get("label", ""),
                    "title": outcome.rendered.get("title", ""),
                    "description": outcome.rendered.get("description", ""),
                    "__fonts": dict(outcome.fonts),
                }
            )
            if not outcome.passed:
                ok = False
        if len(sections) > rows * cols:
            ok = False
        return {"ok": ok, "sections": rendered_sections, "outcomes": outcomes}

    def _section_mode_field_rects(
        self,
        *,
        mode: str,
        geometry: SectionGridGeometry,
        section_index: int,
        rows_override: int | None = None,
        cols_override: int | None = None,
    ) -> dict[str, tuple[float, float, float, float]]:
        mode_spec = SECTION_MODE_SPECS[mode]
        rows = int(rows_override if rows_override is not None else mode_spec["rows"])
        cols = int(cols_override if cols_override is not None else mode_spec["cols"])
        col = section_index % cols
        row = section_index // cols
        cell_w = geometry.content_width_in / cols
        cell_h = geometry.content_height_in / rows
        cell_x = geometry.content_left_in + (col * cell_w)
        cell_y = geometry.content_top_in + (row * cell_h)
        return {
            "label": self._rect_from_ratio(cell_x, cell_y, cell_w, cell_h, geometry.field_ratios.label),
            "title": self._rect_from_ratio(cell_x, cell_y, cell_w, cell_h, geometry.field_ratios.title),
            "description": self._rect_from_ratio(
                cell_x,
                cell_y,
                cell_w,
                cell_h,
                geometry.field_ratios.description,
            ),
        }

    def _rect_from_ratio(
        self,
        x_in: float,
        y_in: float,
        w_in: float,
        h_in: float,
        ratio: tuple[float, float, float, float],
    ) -> tuple[float, float, float, float]:
        rx, ry, rw, rh = ratio
        return (
            x_in + (w_in * rx),
            y_in + (h_in * ry),
            max(0.05, w_in * rw),
            max(0.05, h_in * rh),
        )

    def _evaluate_single_section_fit(
        self,
        *,
        section: dict[str, Any],
        field_rects: Mapping[str, tuple[float, float, float, float]],
        mode_spec: Mapping[str, Any],
        geometry: SectionGridGeometry,
    ) -> SectionFitOutcome:
        reason_codes: list[str] = []
        rendered = {
            "label": str(section.get("label", "") or "").strip(),
            "title": str(section.get("title", "") or "").strip(),
            "description": str(section.get("description", "") or "").strip(),
        }
        fonts: dict[str, float] = {}
        lines: dict[str, int] = {}

        # label
        label_text = rendered["label"]
        label_fit = self._fit_text_for_shape(
            label_text,
            field_rects["label"],
            template_font_pt=float(geometry.template_fonts_pt["label"]),
            min_font_pt=float(mode_spec["min_label_pt"]),
            max_lines=1,
            allow_truncation=False,
            max_truncation_ratio=0.0,
        )
        fonts["label"] = float(label_fit["font_pt"])
        lines["label"] = int(label_fit["lines"])
        if not bool(label_fit["ok"]):
            reason_codes.append("label_overflow")

        # title
        title_text = rendered["title"]
        title_fit = self._fit_text_for_shape(
            title_text,
            field_rects["title"],
            template_font_pt=float(geometry.template_fonts_pt["title"]),
            min_font_pt=float(mode_spec["min_title_pt"]),
            max_lines=int(mode_spec["max_title_lines"]),
            allow_truncation=False,
            max_truncation_ratio=0.0,
        )
        fonts["title"] = float(title_fit["font_pt"])
        lines["title"] = int(title_fit["lines"])
        if not bool(title_fit["ok"]):
            if bool(title_fit.get("line_limit_failed", False)):
                reason_codes.append("title_line_limit")
            else:
                reason_codes.append("title_min_font")

        # description
        desc_text = rendered["description"]
        desc_fit = self._fit_text_for_shape(
            desc_text,
            field_rects["description"],
            template_font_pt=float(geometry.template_fonts_pt["description"]),
            min_font_pt=float(mode_spec["min_desc_pt"]),
            max_lines=int(mode_spec["max_desc_lines"]),
            allow_truncation=bool(mode_spec["allow_desc_truncation"]),
            max_truncation_ratio=float(mode_spec["max_desc_truncation_ratio"]),
        )
        fonts["description"] = float(desc_fit["font_pt"])
        lines["description"] = int(desc_fit["lines"])
        rendered["description"] = str(desc_fit["text"])
        if not bool(desc_fit["ok"]):
            if bool(desc_fit.get("truncation_required", False)):
                reason_codes.append("desc_truncation_required")
            elif bool(desc_fit.get("line_limit_failed", False)):
                reason_codes.append("desc_line_limit")
            else:
                reason_codes.append("desc_min_font")

        return SectionFitOutcome(
            passed=len(reason_codes) == 0,
            section=dict(section),
            rendered=rendered,
            fonts=fonts,
            lines=lines,
            reason_codes=sorted(set(reason_codes)),
        )

    def _fit_text_for_shape(
        self,
        text: str,
        rect: tuple[float, float, float, float],
        *,
        template_font_pt: float,
        min_font_pt: float,
        max_lines: int,
        allow_truncation: bool,
        max_truncation_ratio: float,
    ) -> dict[str, Any]:
        _, _, width_in, height_in = rect
        raw_text = str(text or "").strip()
        if not raw_text:
            return {"ok": True, "text": "", "font_pt": template_font_pt, "lines": 0}

        fit = find_shrink_to_fit_font(
            raw_text,
            width_in=width_in,
            height_in=height_in,
            template_font_pt=template_font_pt,
            min_font_pt=min_font_pt,
            step_pt=0.5,
            line_spacing=1.0,
            vertical_padding_in=0.02,
        )
        if fit is not None:
            chars = estimate_chars_per_line(width_in, fit.font_size_pt)
            wrapped_lines = estimate_multiline_wrapped_lines(raw_text, chars)
            if wrapped_lines <= max_lines:
                return {
                    "ok": True,
                    "text": raw_text,
                    "font_pt": float(fit.font_size_pt),
                    "lines": int(wrapped_lines),
                }
            if not allow_truncation:
                return {
                    "ok": False,
                    "text": raw_text,
                    "font_pt": float(fit.font_size_pt),
                    "lines": int(wrapped_lines),
                    "line_limit_failed": True,
                }
        elif not allow_truncation:
            return {
                "ok": False,
                "text": raw_text,
                "font_pt": float(min_font_pt),
                "lines": max_lines + 1,
            }

        if not allow_truncation:
            return {
                "ok": False,
                "text": raw_text,
                "font_pt": float(min_font_pt),
                "lines": max_lines + 1,
                "truncation_required": True,
            }

        truncated = self._truncate_text_to_fit(
            raw_text,
            width_in=width_in,
            height_in=height_in,
            template_font_pt=template_font_pt,
            min_font_pt=min_font_pt,
            max_lines=max_lines,
            max_truncation_ratio=max_truncation_ratio,
        )
        if truncated is None:
            return {
                "ok": False,
                "text": raw_text,
                "font_pt": float(min_font_pt),
                "lines": max_lines + 1,
                "truncation_required": True,
            }
        return truncated

    def _truncate_text_to_fit(
        self,
        text: str,
        *,
        width_in: float,
        height_in: float,
        template_font_pt: float,
        min_font_pt: float,
        max_lines: int,
        max_truncation_ratio: float,
    ) -> dict[str, Any] | None:
        tokens = str(text or "").split()
        if not tokens:
            return {"ok": True, "text": "", "font_pt": template_font_pt, "lines": 0}
        total_chars = len(str(text))
        min_chars_to_keep = max(1, int(round(total_chars * (1.0 - max_truncation_ratio))))

        current_tokens = list(tokens)
        while current_tokens:
            candidate_text = " ".join(current_tokens).strip()
            if len(candidate_text) < min_chars_to_keep:
                break
            if len(current_tokens) < len(tokens):
                candidate_text = candidate_text.rstrip(".;,:") + "..."
            fit = find_shrink_to_fit_font(
                candidate_text,
                width_in=width_in,
                height_in=height_in,
                template_font_pt=template_font_pt,
                min_font_pt=min_font_pt,
                step_pt=0.5,
                line_spacing=1.0,
                vertical_padding_in=0.02,
            )
            if fit is not None:
                chars = estimate_chars_per_line(width_in, fit.font_size_pt)
                wrapped_lines = estimate_multiline_wrapped_lines(candidate_text, chars)
                if wrapped_lines <= max_lines:
                    return {
                        "ok": True,
                        "text": candidate_text,
                        "font_pt": float(fit.font_size_pt),
                        "lines": int(wrapped_lines),
                    }
            current_tokens = current_tokens[:-1]
        return None

    def _render_section_grid_page(
        self,
        helper: MiniPptxHelper,
        *,
        slide_index: int,
        title_shape_name: str,
        title_text: str,
        sections_mapping: Mapping[str, Any],
        geometry: SectionGridGeometry,
        mode: str,
        page_sections: list[dict[str, Any]],
        forced_page_size: int | None = None,
    ) -> None:
        if title_shape_name:
            title_shape = helper.get_shape_by_name(slide_index, title_shape_name)
            if title_shape is not None:
                helper.replace_text_preserve_format(title_shape, title_text)

        groups = self._materialize_section_shapes(
            helper,
            slide_index=slide_index,
            sections_mapping=sections_mapping,
            needed=max(len(page_sections), 1),
        )
        mode_spec = SECTION_MODE_SPECS.get(mode, SECTION_MODE_SPECS["2x2"])
        capacity = int(mode_spec["capacity"])
        active_count = min(len(page_sections), capacity)
        use_single_section_full_width = active_count == 1

        for index in range(min(active_count, len(groups))):
            group = groups[index]
            section = page_sections[index] if index < len(page_sections) else {}
            if use_single_section_full_width:
                field_rects = self._single_section_full_width_field_rects(
                    geometry=geometry,
                )
            else:
                field_rects = self._section_mode_field_rects(
                    mode=mode,
                    geometry=geometry,
                    section_index=index,
                )
            self._position_and_fill_section_group(
                helper,
                slide_index=slide_index,
                group=group,
                field_rects=field_rects,
                section=section,
                mode=mode,
                apply_title_desc_gap_clamp=use_single_section_full_width,
            )

        # Clear remaining template/cloned groups so nothing stale appears.
        for index in range(active_count, len(groups)):
            group = groups[index]
            for key in ("label_name", "title_name", "description_name"):
                shape = helper.get_shape_by_name(slide_index, str(group.get(key, "")))
                if shape is not None:
                    helper.remove_shape(shape)

    def _single_section_full_width_field_rects(
        self,
        *,
        geometry: SectionGridGeometry,
    ) -> dict[str, tuple[float, float, float, float]]:
        x_in = float(geometry.content_left_in)
        y_in = float(geometry.content_top_in)
        w_in = float(geometry.content_width_in)
        h_in = float(geometry.content_height_in)
        return {
            "label": self._rect_from_ratio(x_in, y_in, w_in, h_in, SECTION_SINGLE_FULLWIDTH_RATIOS["label"]),
            "title": self._rect_from_ratio(x_in, y_in, w_in, h_in, SECTION_SINGLE_FULLWIDTH_RATIOS["title"]),
            "description": self._rect_from_ratio(
                x_in,
                y_in,
                w_in,
                h_in,
                SECTION_SINGLE_FULLWIDTH_RATIOS["description"],
            ),
        }

    def _materialize_section_shapes(
        self,
        helper: MiniPptxHelper,
        *,
        slide_index: int,
        sections_mapping: Mapping[str, Any],
        needed: int,
    ) -> list[dict[str, str]]:
        items = sections_mapping.get("items")
        if not isinstance(items, list):
            return []
        groups: list[dict[str, str]] = []
        for entry in items[:4]:
            if not isinstance(entry, Mapping):
                continue
            groups.append(
                {
                    "label_name": str(entry.get("label", "")).strip(),
                    "title_name": str(entry.get("title", "")).strip(),
                    "description_name": str(entry.get("description", "")).strip(),
                }
            )
        if not groups:
            return []

        while len(groups) < needed:
            source = groups[(len(groups) - 2) % len(groups)] if len(groups) >= 2 else groups[0]
            cloned = {
                "label_name": self._clone_shape_by_name(
                    helper,
                    slide_index=slide_index,
                    source_name=source["label_name"],
                    suffix=f"__clone_{len(groups)+1}",
                ),
                "title_name": self._clone_shape_by_name(
                    helper,
                    slide_index=slide_index,
                    source_name=source["title_name"],
                    suffix=f"__clone_{len(groups)+1}",
                ),
                "description_name": self._clone_shape_by_name(
                    helper,
                    slide_index=slide_index,
                    source_name=source["description_name"],
                    suffix=f"__clone_{len(groups)+1}",
                ),
            }
            groups.append(cloned)
        return groups

    def _clone_shape_by_name(
        self,
        helper: MiniPptxHelper,
        *,
        slide_index: int,
        source_name: str,
        suffix: str,
    ) -> str:
        shape = helper.get_shape_by_name(slide_index, source_name)
        if shape is None:
            return source_name
        slide = helper.get_slide(slide_index)
        new_element = deepcopy(shape._element)
        new_name = f"{source_name}{suffix}"
        c_nv_pr = new_element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
        if c_nv_pr is not None:
            c_nv_pr.set("name", new_name)
        slide.shapes._spTree.insert_element_before(new_element, "p:extLst")
        cloned = slide.shapes[-1]
        return str(getattr(cloned, "name", new_name))

    def _position_and_fill_section_group(
        self,
        helper: MiniPptxHelper,
        *,
        slide_index: int,
        group: Mapping[str, str],
        field_rects: Mapping[str, tuple[float, float, float, float]],
        section: Mapping[str, Any],
        mode: str,
        apply_title_desc_gap_clamp: bool = False,
    ) -> None:
        fonts_raw = section.get("__fonts", {})
        fonts = fonts_raw if isinstance(fonts_raw, Mapping) else {}
        adjusted_rects = dict(field_rects)
        if mode == "2x1" or apply_title_desc_gap_clamp:
            adjusted_rects = self._clamp_section_title_desc_gap(
                field_rects=field_rects,
                gap_in=0.04,
            )
        label_shape = helper.get_shape_by_name(slide_index, str(group.get("label_name", "")))
        title_shape = helper.get_shape_by_name(slide_index, str(group.get("title_name", "")))
        desc_shape = helper.get_shape_by_name(slide_index, str(group.get("description_name", "")))
        for field_name, shape in (
            ("label", label_shape),
            ("title", title_shape),
            ("description", desc_shape),
        ):
            if shape is None:
                continue
            x_in, y_in, w_in, h_in = adjusted_rects[field_name]
            shape.left = int(round(x_in * 914400))
            shape.top = int(round(y_in * 914400))
            shape.width = int(round(w_in * 914400))
            shape.height = int(round(h_in * 914400))
            text = str(section.get(field_name, "") or "").strip()
            helper.replace_text_preserve_format(shape, text)
            font_pt = fonts.get(field_name)
            if font_pt is not None:
                self._set_text_shape_font_size(shape, float(font_pt))
            if field_name == "label":
                self._fit_section_grid_label_single_line(shape, text)

    def _clamp_section_title_desc_gap(
        self,
        *,
        field_rects: Mapping[str, tuple[float, float, float, float]],
        gap_in: float,
    ) -> dict[str, tuple[float, float, float, float]]:
        adjusted = dict(field_rects)
        title = adjusted.get("title")
        desc = adjusted.get("description")
        if title is None or desc is None:
            return adjusted
        tx, ty, tw, th = title
        dx, dy, dw, dh = desc
        min_desc_top = float(ty) + float(th) + max(0.0, float(gap_in))
        if float(dy) < min_desc_top:
            delta = min_desc_top - float(dy)
            dy = min_desc_top
            dh = max(0.05, float(dh) - delta)
        adjusted["description"] = (float(dx), float(dy), float(dw), float(dh))
        return adjusted

    def _fit_section_grid_label_single_line(
        self,
        shape: Any,
        text: str,
        *,
        min_font_pt: float = 8.0,
    ) -> None:
        """Keep section-grid labels on a single line inside their fixed box."""
        if not getattr(shape, "has_text_frame", False):
            return

        label_text = " ".join(str(text or "").split()).strip()
        if not label_text:
            return

        text_frame = shape.text_frame
        text_frame.word_wrap = False

        width_in = emu_to_inches(getattr(shape, "width", 0))
        if width_in <= 0:
            return

        baseline_pt = self._shape_template_font_pt(shape, fallback=18.0)
        start_pt = max(float(min_font_pt), float(baseline_pt))
        required_chars = max(1, len(label_text) + 1)

        chosen_pt = float(min_font_pt)
        candidate = start_pt
        while candidate >= float(min_font_pt) - 1e-6:
            capacity = estimate_chars_per_line(
                width_in,
                candidate,
                average_char_width_factor=0.62,
            )
            if capacity >= required_chars:
                chosen_pt = float(candidate)
                break
            candidate -= 0.5

        self._set_text_shape_font_size(shape, chosen_pt)

    def _render_paginated_table_slides(
        self,
        helper: MiniPptxHelper,
        prototype_index: int,
        content: Mapping[str, Any],
        *,
        mapped_layout: Mapping[str, Any],
        warnings: list[str],
    ) -> int:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return 0

        rows_mapping = placeholders.get("rows")
        if not isinstance(rows_mapping, Mapping):
            return 0

        rows = [
            dict(row)
            for row in content.get("rows", [])
            if isinstance(row, Mapping)
        ]
        table_plan = self._build_table_pagination_plan(
            helper,
            prototype_index,
            rows_mapping,
            rows,
            warnings=warnings,
        )
        pages = table_plan["pages"]
        if not pages:
            pages = [{"rows": [], "row_heights_in": []}]

        rendered = 0
        for page_index, page in enumerate(pages):
            helper.duplicate_slide(prototype_index)
            target_index = helper.slide_count() - 1
            page_content = dict(content)
            page_content["rows"] = page["rows"]
            if page_index < len(pages) - 1:
                page_content["footer_notes"] = []

            self._render_slide_content(
                helper,
                target_index,
                page_content,
                mapped_layout=mapped_layout,
                layout_id="info_table",
            )
            self._apply_mapped_title_shrink_to_fit(
                helper,
                target_index,
                content=page_content,
                mapped_layout=mapped_layout,
                fallback_template_pt=40.0,
                min_font_pt=12.0,
            )
            self._apply_table_page_geometry(
                helper,
                target_index,
                rows_mapping,
                page["active_row_heights_in"],
                float(table_plan["font_size_pt"]),
                target_table_height_in=float(page.get("target_table_height_in", 0.0)),
            )
            rendered += 1
        return rendered

    def _build_table_pagination_plan(
        self,
        helper: MiniPptxHelper,
        prototype_index: int,
        table_mapping: Mapping[str, Any],
        rows: list[dict[str, Any]],
        *,
        warnings: list[str],
    ) -> dict[str, Any]:
        shape = helper.get_shape_by_name(prototype_index, str(table_mapping.get("name", "")))
        if shape is None or not getattr(shape, "has_table", False):
            warnings.append("Could not paginate info_table because the table placeholder was not found.")
            return {
                "font_size_pt": 10.0,
                "pages": [
                    {
                        "rows": rows,
                        "active_row_heights_in": [],
                        "used_body_height_in": 0.0,
                        "is_final_page": True,
                        "target_table_height_in": 0.0,
                    }
                ],
            }

        table = shape.table
        start_row = int(table_mapping.get("start_row", 0) or 0)
        physical_row_count = max(1, len(table.rows) - start_row)
        table_height_in = emu_to_inches(shape.height)
        fixed_pre_start_height_in = self._table_rows_height_in(table, 0, start_row)
        body_height_in = max(0.01, table_height_in - fixed_pre_start_height_in)
        columns = table_mapping.get("columns", {})
        columns = columns if isinstance(columns, Mapping) else {}
        column_widths_in = {
            str(field): emu_to_inches(table.columns[int(column_index)].width) * 0.92
            for field, column_index in columns.items()
            if 0 <= int(column_index) < len(table.columns)
        }

        pagination = table_mapping.get("pagination", {})
        pagination = pagination if isinstance(pagination, Mapping) else {}
        template_font_pt = float(
            pagination.get("template_font_pt")
            or self._infer_table_template_font_size(table, start_row, columns)
            or 10.0
        )
        min_font_pt = float(pagination.get("min_font_pt", 7.0))
        step_pt = float(pagination.get("step_pt", 0.5))
        line_spacing = float(pagination.get("line_spacing", 1.0))
        vertical_padding_in = float(pagination.get("vertical_padding_in", 0.08))
        min_row_height_in = float(
            pagination.get(
                "min_row_height_in",
                max(0.22, estimate_line_height_in(template_font_pt) + 0.06),
            )
        )

        font_size_pt = self._choose_table_font_size(
            rows,
            column_widths_in,
            body_height_in,
            template_font_pt=template_font_pt,
            min_font_pt=min_font_pt,
            step_pt=step_pt,
            line_spacing=line_spacing,
            vertical_padding_in=vertical_padding_in,
        )
        row_heights = self._estimate_table_row_heights(
            rows,
            column_widths_in,
            font_size_pt,
            line_spacing=line_spacing,
            vertical_padding_in=vertical_padding_in,
            min_row_height_in=min_row_height_in,
        )
        pages = self._paginate_table_rows(
            rows,
            row_heights,
            body_capacity_height_in=body_height_in,
            max_rows_per_page=physical_row_count,
            fixed_pre_start_height_in=fixed_pre_start_height_in,
            full_table_height_in=table_height_in,
        )

        return {
            "font_size_pt": font_size_pt,
            "pages": pages,
        }

    def _estimate_table_row_heights(
        self,
        rows: list[dict[str, Any]],
        column_widths_in: Mapping[str, float],
        font_size_pt: float,
        *,
        line_spacing: float,
        vertical_padding_in: float,
        min_row_height_in: float,
    ) -> list[float]:
        return [
            self._estimate_table_row_height(
                row,
                column_widths_in,
                font_size_pt,
                line_spacing=line_spacing,
                vertical_padding_in=vertical_padding_in,
                min_row_height_in=min_row_height_in,
            )
            for row in rows
        ]

    def _paginate_table_rows(
        self,
        rows: list[dict[str, Any]],
        row_heights_in: list[float],
        *,
        body_capacity_height_in: float,
        max_rows_per_page: int,
        fixed_pre_start_height_in: float,
        full_table_height_in: float,
    ) -> list[dict[str, Any]]:
        # Leave a small visual buffer so wrapped rows do not crowd the section below.
        safety_margin_in = 0.04
        capacity = max(0.01, float(body_capacity_height_in) - safety_margin_in)
        max_rows = max(1, int(max_rows_per_page))
        pages: list[dict[str, Any]] = []
        cursor = 0

        while cursor < len(rows):
            used_height = 0.0
            page_rows: list[dict[str, Any]] = []
            active_heights: list[float] = []

            while cursor < len(rows) and len(page_rows) < max_rows:
                next_height = min(max(0.01, float(row_heights_in[cursor])), capacity)
                if page_rows and used_height + next_height > capacity + 1e-6:
                    break
                page_rows.append(rows[cursor])
                active_heights.append(next_height)
                used_height += next_height
                cursor += 1

                if used_height >= capacity - 1e-6:
                    break

            if not page_rows:
                next_height = min(max(0.01, float(row_heights_in[cursor])), capacity)
                page_rows.append(rows[cursor])
                active_heights.append(next_height)
                used_height = next_height
                cursor += 1

            pages.append(
                {
                    "rows": page_rows,
                    "active_row_heights_in": active_heights,
                    "used_body_height_in": used_height,
                }
            )

        if not pages:
            pages = [
                {
                    "rows": [],
                    "active_row_heights_in": [],
                    "used_body_height_in": 0.0,
                }
            ]

        for index, page in enumerate(pages):
            is_final_page = index == len(pages) - 1
            used_height = max(0.0, float(page["used_body_height_in"]))
            active_count = len(page["rows"])
            page["is_final_page"] = is_final_page
            page["target_table_height_in"] = (
                fixed_pre_start_height_in + used_height
                if active_count < max_rows
                else float(full_table_height_in)
            )

        return pages

    def _table_rows_height_in(self, table: Any, start_row: int, end_row: int) -> float:
        bounded_start = max(0, int(start_row))
        bounded_end = min(max(0, int(end_row)), len(table.rows))
        if bounded_end <= bounded_start:
            return 0.0
        return sum(
            emu_to_inches(table.rows[row_index].height)
            for row_index in range(bounded_start, bounded_end)
        )

    def _infer_table_template_font_size(
        self,
        table: Any,
        start_row: int,
        columns: Mapping[str, Any],
    ) -> float | None:
        sizes: list[float] = []
        for row_index in range(start_row, len(table.rows)):
            for column_index in columns.values():
                try:
                    cell = table.cell(row_index, int(column_index))
                except Exception:
                    continue
                sizes.extend(self._cell_font_sizes(cell))
        if not sizes:
            return None
        return min(sizes)

    def _cell_font_sizes(self, cell: Any) -> list[float]:
        sizes: list[float] = []
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                size = getattr(run.font, "size", None)
                if size is not None:
                    sizes.append(float(size.pt))
        return sizes

    def _choose_table_font_size(
        self,
        rows: list[dict[str, Any]],
        column_widths_in: Mapping[str, float],
        table_height_in: float,
        *,
        template_font_pt: float,
        min_font_pt: float,
        step_pt: float,
        line_spacing: float,
        vertical_padding_in: float,
    ) -> float:
        chosen = float(template_font_pt)
        for row in rows:
            for field_name, width_in in column_widths_in.items():
                text = self._format_placeholder_value(row.get(field_name, ""))
                fit = find_shrink_to_fit_font(
                    text,
                    width_in=width_in,
                    height_in=table_height_in,
                    template_font_pt=template_font_pt,
                    min_font_pt=min_font_pt,
                    step_pt=step_pt,
                    line_spacing=line_spacing,
                    vertical_padding_in=vertical_padding_in,
                )
                if fit is None:
                    chosen = min(chosen, min_font_pt)
                else:
                    chosen = min(chosen, float(fit.font_size_pt))
        return chosen

    def _estimate_table_row_height(
        self,
        row: Mapping[str, Any],
        column_widths_in: Mapping[str, float],
        font_size_pt: float,
        *,
        line_spacing: float,
        vertical_padding_in: float,
        min_row_height_in: float,
    ) -> float:
        heights = [float(min_row_height_in)]
        for field_name, width_in in column_widths_in.items():
            text = self._format_placeholder_value(row.get(field_name, ""))
            heights.append(
                estimate_text_height_in(
                    text,
                    width_in=width_in,
                    font_size_pt=font_size_pt,
                    line_spacing=line_spacing,
                )
                + vertical_padding_in
            )
        return max(heights)

    def _apply_table_page_geometry(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        table_mapping: Mapping[str, Any],
        active_row_heights_in: list[float],
        font_size_pt: float,
        *,
        target_table_height_in: float,
    ) -> None:
        shape = helper.get_shape_by_name(slide_index, str(table_mapping.get("name", "")))
        if shape is None or not getattr(shape, "has_table", False):
            return

        table = shape.table
        template_table_height_in = emu_to_inches(shape.height)
        start_row = int(table_mapping.get("start_row", 0) or 0)
        original_row_count = len(table.rows)
        used_count = len(active_row_heights_in)

        for offset, row_height_in in enumerate(active_row_heights_in):
            row_index = start_row + offset
            if row_index >= original_row_count:
                break
            table.rows[row_index].height = int(round(row_height_in * 914400))
            for column_index in range(len(table.columns)):
                cell = table.cell(row_index, column_index)
                self._set_cell_font_size(cell, font_size_pt)

        self._trim_table_rows(table, keep_row_count=start_row + used_count)

        if float(target_table_height_in) > 0:
            safety_margin_in = 0.04
            max_allowed_height_in = max(0.01, template_table_height_in - safety_margin_in)
            final_height_in = min(float(target_table_height_in), max_allowed_height_in)
            shape.height = int(round(max(0.01, final_height_in) * 914400))

    def _trim_table_rows(self, table: Any, *, keep_row_count: int) -> None:
        tbl = getattr(table, "_tbl", None)
        if tbl is None:
            return
        keep_count = max(0, int(keep_row_count))
        row_elements = list(getattr(tbl, "tr_lst", []))
        for row_element in reversed(row_elements[keep_count:]):
            tbl.remove(row_element)

    def _set_cell_font_size(self, cell: Any, font_size_pt: float) -> None:
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(font_size_pt)

    def _set_text_shape_font_size(self, shape: Any, font_size_pt: float) -> None:
        if not getattr(shape, "has_text_frame", False):
            return
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(font_size_pt)

    def _set_text_shape_bold(self, shape: Any, is_bold: bool) -> None:
        if not getattr(shape, "has_text_frame", False):
            return
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.bold = bool(is_bold)

    def _select_content_source(
        self,
        slide_entry: Mapping[str, Any],
        *,
        prefer_input: bool = False,
    ) -> dict[str, Any]:
        keys = (
            ("input_content", "normalized_content", "validated_content", "content")
            if prefer_input
            else ("normalized_content", "validated_content", "input_content", "content")
        )
        for key in keys:
            value = slide_entry.get(key)
            if isinstance(value, Mapping):
                return {str(k): v for k, v in dict(value).items()}
        return {}

    def _render_slide_content(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        content: Mapping[str, Any],
        *,
        mapped_layout: Mapping[str, Any] | None = None,
        layout_id: str = "",
    ) -> None:
        placeholders = {}
        if isinstance(mapped_layout, Mapping):
            raw_placeholders = mapped_layout.get("placeholders")
            if isinstance(raw_placeholders, Mapping):
                placeholders = {str(k): v for k, v in dict(raw_placeholders).items()}

        if placeholders and self._render_placeholders(
            helper,
            slide_index,
            content,
            placeholders,
            layout_id=layout_id,
        ):
            return

        text_shapes = helper.sort_shapes_reading_order(helper.text_shapes(slide_index))
        if not text_shapes:
            return

        blocks = self._build_text_blocks(content)
        if not blocks:
            helper.clear_text_shapes(slide_index)
            return

        helper.replace_text_preserve_format(text_shapes[0], blocks[0])
        remaining_blocks = blocks[1:]
        if not remaining_blocks:
            for shape in text_shapes[1:]:
                helper.replace_text_preserve_format(shape, "")
            return

        target_shapes = text_shapes[1:] or text_shapes[:1]
        for idx, shape in enumerate(target_shapes):
            if idx < len(remaining_blocks):
                helper.replace_text_preserve_format(shape, remaining_blocks[idx])
            elif idx == len(target_shapes) - 1 and len(remaining_blocks) > len(target_shapes):
                overflow_start = len(target_shapes) - 1
                helper.replace_text_preserve_format(
                    shape,
                    "\n\n".join(remaining_blocks[overflow_start:]),
                )
            else:
                helper.replace_text_preserve_format(shape, "")

    def _render_placeholders(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        content: Mapping[str, Any],
        placeholders: Mapping[str, Any],
        *,
        layout_id: str = "",
    ) -> bool:
        rendered_any = False
        for field, mapping in placeholders.items():
            if not isinstance(mapping, Mapping):
                continue
            if "contains" in mapping:
                rendered_any = self._render_composite_text_placeholder(
                    helper,
                    slide_index,
                    mapping,
                    content,
                    layout_id=layout_id,
                ) or rendered_any
                continue
            field_name = str(field)
            value = content.get(field_name)
            mapping_type = str(mapping.get("type", "")).strip()

            if mapping_type == "table":
                rendered_any = self._render_table_placeholder(
                    helper,
                    slide_index,
                    mapping,
                    value,
                ) or rendered_any
                continue

            if mapping_type == "repeated_group":
                rendered_any = self._render_repeated_group_placeholder(
                    helper,
                    slide_index,
                    mapping,
                    value,
                ) or rendered_any
                continue

            if "names" in mapping:
                rendered_any = self._render_multi_shape_placeholder(
                    helper,
                    slide_index,
                    mapping,
                    value,
                ) or rendered_any
                continue

            shape_name = str(mapping.get("name", "")).strip()
            shape = helper.get_shape_by_name(slide_index, shape_name)
            if shape is None:
                continue
            helper.replace_text_preserve_format(shape, self._format_placeholder_value(value))
            rendered_any = True
        return rendered_any

    def _render_composite_text_placeholder(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        mapping: Mapping[str, Any],
        content: Mapping[str, Any],
        *,
        layout_id: str = "",
    ) -> bool:
        contains = mapping.get("contains")
        if not isinstance(contains, list):
            return False

        shape_name = str(mapping.get("name", "")).strip()
        shape = helper.get_shape_by_name(slide_index, shape_name)
        if shape is None or not getattr(shape, "has_text_frame", False):
            return False

        text_frame = shape.text_frame
        text_frame.clear()
        wrote_any = False

        for field in contains:
            field_name = str(field or "").strip()
            if not field_name:
                continue
            value = content.get(field_name)
            if value is None:
                continue

            should_bullet = isinstance(value, list) or self._is_bullet_field_type(
                layout_id=layout_id,
                field_name=field_name,
            )
            if should_bullet:
                for item in self._as_list(value):
                    paragraph = self._next_text_frame_paragraph(text_frame, wrote_any)
                    self._set_paragraph_text(paragraph, item)
                    self._set_paragraph_bullet(paragraph, enabled=True)
                    wrote_any = True
                continue

            text = self._as_text(value)
            if not text:
                continue
            paragraph = self._next_text_frame_paragraph(text_frame, wrote_any)
            self._set_paragraph_text(paragraph, text)
            self._set_paragraph_bullet(paragraph, enabled=False)
            wrote_any = True

        if not wrote_any:
            first = text_frame.paragraphs[0]
            self._set_paragraph_text(first, "")
            self._set_paragraph_bullet(first, enabled=False)
        return wrote_any

    def _is_bullet_field_type(self, *, layout_id: str, field_name: str) -> bool:
        contract = self.layout_registry.get(layout_id)
        if contract is None:
            return False
        return str(contract.field_types.get(str(field_name), "")).strip() == "bullet_list"

    @staticmethod
    def _next_text_frame_paragraph(text_frame: Any, wrote_any: bool):
        if not wrote_any and text_frame.paragraphs:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()
        return paragraph

    def _set_paragraph_text(self, paragraph: Any, text: str) -> None:
        text_value = str(text or "")
        if paragraph.runs:
            for run in paragraph.runs:
                run.text = ""
            paragraph.runs[0].text = text_value
            return
        paragraph.text = text_value

    def _set_paragraph_bullet(self, paragraph: Any, *, enabled: bool) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        for child in list(p_pr):
            local_name = str(child.tag).split("}")[-1]
            if local_name in {"buNone", "buAutoNum", "buChar", "buBlip"}:
                p_pr.remove(child)
        if enabled:
            paragraph.level = 0
            bu_char = OxmlElement("a:buChar")
            bu_char.set("char", "•")
            p_pr.append(bu_char)
        else:
            bu_none = OxmlElement("a:buNone")
            p_pr.append(bu_none)

    def _render_multi_shape_placeholder(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        mapping: Mapping[str, Any],
        value: Any,
    ) -> bool:
        names = mapping.get("names")
        if not isinstance(names, list):
            return False
        values = self._as_list(value)
        rendered_any = False
        for idx, name in enumerate(names):
            shape = helper.get_shape_by_name(slide_index, str(name))
            if shape is None:
                continue
            text = values[idx] if idx < len(values) else ""
            helper.replace_text_preserve_format(shape, text)
            rendered_any = True
        return rendered_any

    def _render_table_placeholder(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        mapping: Mapping[str, Any],
        value: Any,
    ) -> bool:
        shape = helper.get_shape_by_name(slide_index, str(mapping.get("name", "")))
        if shape is None or not getattr(shape, "has_table", False):
            return False
        rows = value if isinstance(value, list) else []
        table = shape.table
        start_row = int(mapping.get("start_row", 0) or 0)
        columns = mapping.get("columns", {})
        columns = columns if isinstance(columns, Mapping) else {}

        for row_index in range(start_row, len(table.rows)):
            item_index = row_index - start_row
            item = rows[item_index] if item_index < len(rows) and isinstance(rows[item_index], Mapping) else {}
            for field_name, column_index in columns.items():
                try:
                    cell = table.cell(row_index, int(column_index))
                except Exception:
                    continue
                cell.text = self._format_placeholder_value(item.get(str(field_name), ""))
        return True

    def _render_repeated_group_placeholder(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        mapping: Mapping[str, Any],
        value: Any,
    ) -> bool:
        item_mappings = mapping.get("items")
        if not isinstance(item_mappings, list):
            return False
        values = value if isinstance(value, list) else []
        rendered_any = False
        for idx, item_mapping in enumerate(item_mappings):
            if not isinstance(item_mapping, Mapping):
                continue
            item = values[idx] if idx < len(values) and isinstance(values[idx], Mapping) else {}
            for field_name, shape_name in item_mapping.items():
                shape = helper.get_shape_by_name(slide_index, str(shape_name))
                if shape is None:
                    continue
                helper.replace_text_preserve_format(
                    shape,
                    self._format_placeholder_value(item.get(str(field_name), "")),
                )
                rendered_any = True
        return rendered_any

    def _build_text_blocks(self, content: Mapping[str, Any]) -> list[str]:
        blocks: list[str] = []
        used_fields: set[str] = set()

        title = self._as_text(content.get("title"))
        if title:
            blocks.append(title)
            used_fields.add("title")

        for key in ("intro_paragraphs",):
            values = self._as_list(content.get(key))
            if values:
                blocks.append("\n".join(values))
                used_fields.add(key)

        pair_definitions = [
            ("card_titles", "card_descriptions"),
            ("point_titles", "point_descriptions"),
            ("step_titles", "step_descriptions"),
        ]
        for titles_key, desc_key in pair_definitions:
            titles = self._as_list(content.get(titles_key))
            descs = self._as_list(content.get(desc_key))
            if not titles and not descs:
                continue
            used_fields.update({titles_key, desc_key})
            max_len = max(len(titles), len(descs))
            for idx in range(max_len):
                title_text = titles[idx] if idx < len(titles) else ""
                desc_text = descs[idx] if idx < len(descs) else ""
                if title_text and desc_text:
                    blocks.append(f"{idx + 1}. {title_text}\n{desc_text}")
                elif title_text:
                    blocks.append(f"{idx + 1}. {title_text}")
                elif desc_text:
                    blocks.append(f"{idx + 1}. {desc_text}")

        for key, value in content.items():
            field = str(key)
            if field in used_fields or field == "image":
                continue
            if isinstance(value, list):
                values = self._as_list(value)
                if values:
                    blocks.append("\n".join(values))
                continue
            text = self._as_text(value)
            if text:
                blocks.append(text)

        return blocks

    def _replace_image_for_mapped_placeholder(
        self,
        helper: MiniPptxHelper,
        slide_index: int,
        *,
        mapped_layout: Mapping[str, Any],
        image_path: Path,
    ) -> bool:
        placeholders = mapped_layout.get("placeholders")
        if not isinstance(placeholders, Mapping):
            return False
        image_mapping = placeholders.get("image")
        if not isinstance(image_mapping, Mapping):
            image_mapping = placeholders.get("media")
        if not isinstance(image_mapping, Mapping):
            return False

        shape_name = str(image_mapping.get("name", "")).strip()
        if not shape_name:
            return False

        shape = helper.get_shape_by_name(slide_index, shape_name)
        if shape is None:
            return False

        fit_mode = str(image_mapping.get("fit_mode", "stretch")).strip().lower()
        final_image_path = image_path
        if fit_mode and fit_mode != "stretch":
            try:
                prepared_path = prepare_image_for_fit_mode(
                    str(image_path),
                    target_width_emu=int(shape.width),
                    target_height_emu=int(shape.height),
                    fit_mode=fit_mode,
                )
                candidate = Path(prepared_path)
                if candidate.is_file():
                    final_image_path = candidate
            except Exception:
                final_image_path = image_path

        return bool(helper.replace_picture_shape(shape, final_image_path))

    def _extract_image_ref(self, content: Mapping[str, Any]) -> str | None:
        image_text = self._as_text(content.get("image"))
        if image_text:
            return image_text
        media_text = self._as_text(content.get("media"))
        return media_text or None

    def _resolve_image_path(self, image_ref: str) -> Path | None:
        ref = str(image_ref).strip()
        if not ref:
            return None

        raw = Path(ref)
        if raw.is_file():
            return raw.resolve()

        package_root = Path(__file__).resolve().parents[1]
        search_dirs = [
            package_root / "examples" / "media",
            package_root / "examples" / "images",
            Path.cwd() / "data" / "1" / "images",
            Path.cwd() / "data" / "1" / "media",
        ]

        candidate_names = [raw.name]
        if raw.suffix == "":
            for ext in (".png", ".jpg", ".jpeg", ".webp"):
                candidate_names.append(f"{raw.name}{ext}")

        for directory in search_dirs:
            if not directory.exists():
                continue
            for name in candidate_names:
                path = directory / name
                if path.is_file():
                    return path.resolve()
        return None

    @staticmethod
    def _as_text(value: Any) -> str:
        if value is None:
            return ""
        text = str(value).strip()
        return text

    @staticmethod
    def _as_list(value: Any) -> list[str]:
        if value is None:
            return []
        if isinstance(value, list):
            return [str(item).strip() for item in value if str(item).strip()]
        text = str(value).strip()
        if not text:
            return []
        return [part.strip() for part in text.splitlines() if part.strip()]

    def _format_placeholder_value(self, value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, list):
            return "\n".join(self._as_list(value))
        if isinstance(value, Mapping):
            return "\n".join(
                self._as_text(v)
                for v in value.values()
                if self._as_text(v)
            )
        return self._as_text(value)


def load_json_payload(path: str | Path) -> Any:
    file_path = Path(path)
    return json.loads(file_path.read_text(encoding="utf-8"))


def normalize_request_payload(payload: Any) -> dict[str, Any]:
    if isinstance(payload, list):
        if not payload:
            raise ValueError("Request payload list is empty.")
        first = payload[0]
        if not isinstance(first, Mapping):
            raise ValueError("First request item must be an object.")
        return dict(first)
    if isinstance(payload, Mapping):
        return dict(payload)
    raise ValueError("Request payload must be an object or list of objects.")
