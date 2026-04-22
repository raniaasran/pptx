from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


class MiniPptxHelper:
    def __init__(self, template_path: str | Path):
        self.template_path = str(template_path)
        self.prs = Presentation(self.template_path)

    def slide_count(self) -> int:
        return len(self.prs.slides)

    def get_slide(self, slide_index: int):
        return self.prs.slides[slide_index]

    def text_shapes(self, slide_index: int) -> list:
        slide = self.get_slide(slide_index)
        return [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]

    def picture_shapes(self, slide_index: int) -> list:
        slide = self.get_slide(slide_index)
        return [
            shape
            for shape in slide.shapes
            if int(getattr(shape, "shape_type", 0)) == int(MSO_SHAPE_TYPE.PICTURE)
        ]

    def replace_text_preserve_format(self, shape, new_text: str) -> bool:
        if not getattr(shape, "has_text_frame", False):
            return False

        text_frame = shape.text_frame
        first_run = None
        if text_frame.paragraphs and text_frame.paragraphs[0].runs:
            first_run = text_frame.paragraphs[0].runs[0]

        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                run.text = ""

        if first_run is not None:
            first_run.text = str(new_text or "")
        else:
            text_frame.text = str(new_text or "")
        return True

    def clear_text_shapes(self, slide_index: int) -> None:
        for shape in self.text_shapes(slide_index):
            self.replace_text_preserve_format(shape, "")

    def replace_picture_shape(self, shape, image_path: str | Path) -> bool:
        path = Path(image_path)
        if not path.is_file():
            return False

        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        slide = shape.part.slide
        element = shape._element
        element.getparent().remove(element)
        slide.shapes.add_picture(str(path), left, top, width, height)
        return True

    def delete_slide(self, slide_index: int) -> None:
        slides = list(self.prs.slides._sldIdLst)
        slide_id_to_remove = slides[slide_index].rId
        self.prs.part.drop_rel(slide_id_to_remove)
        self.prs.slides._sldIdLst.remove(slides[slide_index])

    def duplicate_slide(self, slide_index: int, insert_after_index: int | None = None):
        src_slide = self.prs.slides[slide_index]
        layout = src_slide.slide_layout
        new_slide = self.prs.slides.add_slide(layout)

        if insert_after_index is not None:
            sld_id_list = self.prs.slides._sldIdLst
            sld_id_elems = list(sld_id_list)
            new_elem = sld_id_elems[-1]
            sld_id_list.remove(new_elem)
            sld_id_list.insert(insert_after_index + 1, new_elem)

        rel_map: dict[str, str] = {}
        for rel_id in list(src_slide.part.rels):
            rel_obj = src_slide.part.rels[rel_id]
            try:
                target_part = rel_obj.target_part
            except AttributeError:
                continue
            try:
                new_rel_id = new_slide.part.relate_to(target_part, rel_obj.reltype)
            except Exception:
                continue
            rel_map[rel_id] = new_rel_id

        for shape in src_slide.shapes:
            new_element = deepcopy(shape._element)
            for blip in new_element.findall(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
            ):
                rid = blip.get(qn("r:embed"))
                if rid and rid in rel_map:
                    blip.set(qn("r:embed"), rel_map[rid])

            for node in new_element.findall(".//*"):
                for attr_name in ("r:embed", "r:id"):
                    q_attr = qn(attr_name)
                    value = node.get(q_attr)
                    if value and value in rel_map:
                        node.set(q_attr, rel_map[value])

            new_slide.shapes._spTree.insert_element_before(new_element, "p:extLst")

        return new_slide

    def save(self, output_path: str | Path) -> None:
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out))

    @staticmethod
    def sort_shapes_reading_order(shapes: Iterable) -> list:
        return sorted(shapes, key=lambda shape: (int(shape.top), int(shape.left)))
