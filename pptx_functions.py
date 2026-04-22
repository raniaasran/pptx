from pathlib import Path
import math
import re, json, os
from copy import deepcopy
from collections import OrderedDict
import base64

from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Pt

from helpers.config.general_config import Settings
from helpers.project_paths import ProjectPaths
from services.content_services.slide_payloads import (
    TEMPLATE_KEY_FIELD,
    canonicalize_slide_payload,
    strip_slide_payload_metadata,
)

config = Settings()

EMU_PER_INCH = 914400

_BULLET_KEY_PRIORITY = {
    "first_bullet": 0,
    "second_bullet": 1,
    "third_bullet": 2,
}

_DENSITY_RULES = {
    "sparse": {
        "font_size_pt": 30,
        "min_font_pt": 23,
        "line_spacing": 1.10,
        "space_after_pt": 2,
    },
    "normal": {
        "font_size_pt": 28,
        "min_font_pt": 22,
        "line_spacing": 1.08,
        "space_after_pt": 2,
    },
    "dense": {
        "font_size_pt": 24,
        "min_font_pt": 20,
        "line_spacing": 1.04,
        "space_after_pt": 1,
    },
    "overflow": {
        "font_size_pt": 22,
        "min_font_pt": 18,
        "line_spacing": 1.02,
        "space_after_pt": 1,
    },
}

_MIN_READABLE_FONT_PT = 18
_MIN_TITLE_FONT_PT = 22
_BASE_FIT_FONT_PT = 21
_BASE_CHARS_PER_LINE = 72
_TEMPLATE_KEY_META_FIELD = "__template_key__"

_LAYOUT_STAGES = (
    {
        "name": "expand_area",
        "top_raise_in": 0.42,
        "second_push_down_in": 0.60,
        "second_height_reduce_in": 0.20,
        "second_min_height_in": 1.45,
        "below_gap_in": 0.08,
        "item_gap_in": 0.09,
        "line_spacing_factor": 1.00,
        "width_expand_in": 0.08,
        "marker_scale": 1.00,
        "decorative_scale": 1.00,
    },
    {
        "name": "decorative_reduction",
        "top_raise_in": 0.58,
        "second_push_down_in": 0.92,
        "second_height_reduce_in": 0.42,
        "second_min_height_in": 1.18,
        "below_gap_in": 0.06,
        "item_gap_in": 0.06,
        "line_spacing_factor": 0.97,
        "width_expand_in": 0.24,
        "marker_scale": 0.88,
        "decorative_scale": 0.90,
    },
)

_STORY_IMG_FAMILY_REQUIRED_KEYS = {
    "title_shape",
    "first_textbox",
    "second_textbox",
    "third_textbox",
    "fourth_textbox",
    "fifth_textbox",
    "picture_placeholder",
}
_STORY_IMG_FLOW_FIELDS = (
    "first_textbox",
    "second_textbox",
    "third_textbox",
    "fourth_textbox",
    "fifth_textbox",
)
_STORY_IMG_MIN_HEIGHTS = {
    "first_textbox": 1.00,
    "second_textbox": 0.74,
    "third_textbox": 0.52,
    "fourth_textbox": 0.52,
    "fifth_textbox": 0.40,
}
_STORY_IMG_LAYOUT_STAGES = (
    {
        "name": "expand_text_region",
        "column_shift_left_with_image_in": 0.32,
        "column_shift_left_without_image_in": 1.25,
        "column_expand_right_in": 0.40,
        "top_raise_in": 0.24,
        "bottom_extend_in": 0.58,
        "gap_scale": 0.78,
        "min_gap_in": 0.05,
        "line_spacing_factor": 1.00,
        "decorative_scale": 1.00,
    },
    {
        "name": "reduce_nonessential",
        "column_shift_left_with_image_in": 0.52,
        "column_shift_left_without_image_in": 1.65,
        "column_expand_right_in": 0.74,
        "top_raise_in": 0.40,
        "bottom_extend_in": 0.90,
        "gap_scale": 0.62,
        "min_gap_in": 0.04,
        "line_spacing_factor": 0.97,
        "decorative_scale": 0.90,
    },
)

_IMG_FOCUS_REPEAT_TEMPLATE_KEY = "text_with_image_focus_repeat"
_IMG_FOCUS_REPEAT_REQUIRED_KEYS = {
    "title_shape",
    "first_textbox",
    "second_textbox",
    "picture_placeholder",
}
_IMG_FOCUS_REPEAT_FLOW_FIELDS = (
    "second_textbox",
    "first_textbox",
)
_IMG_FOCUS_REPEAT_MIN_HEIGHTS = {
    "second_textbox": 1.40,
    "first_textbox": 0.72,
}
_IMG_FOCUS_REPEAT_LAYOUT_STAGES = (
    {
        "name": "expand_text_region",
        "column_shift_left_in": 0.24,
        "column_expand_right_in": 0.34,
        "column_no_image_bonus_right_in": 0.92,
        "image_text_gap_in": 0.36,
        "top_raise_in": 0.20,
        "bottom_extend_in": 0.56,
        "below_image_extend_in": 0.14,
        "gap_scale": 0.86,
        "min_gap_in": 0.06,
        "line_spacing_factor": 1.00,
        "decorative_scale": 1.00,
        "allow_shrink": False,
    },
    {
        "name": "reduce_nonessential",
        "column_shift_left_in": 0.42,
        "column_expand_right_in": 0.56,
        "column_no_image_bonus_right_in": 1.16,
        "image_text_gap_in": 0.30,
        "top_raise_in": 0.34,
        "bottom_extend_in": 0.86,
        "below_image_extend_in": 0.24,
        "gap_scale": 0.72,
        "min_gap_in": 0.05,
        "line_spacing_factor": 0.98,
        "decorative_scale": 0.90,
        "allow_shrink": False,
    },
    {
        "name": "bounded_font_reduction",
        "column_shift_left_in": 0.46,
        "column_expand_right_in": 0.60,
        "column_no_image_bonus_right_in": 1.22,
        "image_text_gap_in": 0.28,
        "top_raise_in": 0.38,
        "bottom_extend_in": 0.96,
        "below_image_extend_in": 0.28,
        "gap_scale": 0.68,
        "min_gap_in": 0.05,
        "line_spacing_factor": 0.97,
        "decorative_scale": 0.88,
        "allow_shrink": True,
    },
)


def _to_emu(value_in: float) -> int:
    return int(round(float(value_in) * EMU_PER_INCH))


def _to_in(value_emu: int) -> float:
    return float(value_emu) / EMU_PER_INCH


def _ordered_bullet_keys(slide_payload: dict) -> list[str]:
    bullet_keys = [
        str(key)
        for key in slide_payload.keys()
        if "bullet" in str(key).lower()
    ]
    return sorted(
        bullet_keys,
        key=lambda key: (_BULLET_KEY_PRIORITY.get(str(key).lower(), 999), str(key)),
    )


def _split_bullet_text(value: str) -> list[str]:
    raw = str(value or "").strip()
    if not raw:
        return []

    cleaned = re.sub(r"\r\n?", "\n", raw)
    tokens: list[str] = []
    for line in cleaned.split("\n"):
        piece = str(line).strip()
        if not piece:
            continue
        for part in re.split(r"\s*[;|]\s*", piece):
            normalized = re.sub(r"^\s*[-*\u2022]+\s*", "", str(part).strip())
            if normalized:
                tokens.append(normalized)
    return tokens


def _collect_bullet_items(slide_payload: dict, bullet_keys: list[str]) -> list[str]:
    items: list[str] = []
    for key in bullet_keys:
        items.extend(_split_bullet_text(slide_payload.get(key, "")))
    return items


def _estimate_chars_per_line(width_in: float, font_size_pt: float) -> int:
    width_factor = max(0.65, float(width_in) / 14.9)
    font_factor = _BASE_FIT_FONT_PT / max(12.0, float(font_size_pt))
    return max(16, int(round(_BASE_CHARS_PER_LINE * width_factor * font_factor)))


def _line_height_in(font_size_pt: float, line_spacing: float) -> float:
    return (float(font_size_pt) / 72.0) * float(line_spacing) * 1.22


def _estimate_wrapped_lines(text: str, chars_per_line: int = _BASE_CHARS_PER_LINE) -> int:
    compact = re.sub(r"\s+", " ", str(text or "").strip())
    if not compact:
        return 0
    return max(1, math.ceil(len(compact) / max(1, chars_per_line)))


def _estimate_total_bullet_lines(
    items: list[str],
    chars_per_line: int = _BASE_CHARS_PER_LINE,
) -> int:
    return sum(_estimate_wrapped_lines(item, chars_per_line=chars_per_line) for item in items)


def _classify_bullet_density(
    items: list[str],
    chars_per_line: int = _BASE_CHARS_PER_LINE,
) -> dict[str, int | str]:
    bullet_count = len(items)
    total_chars = sum(len(item) for item in items)
    estimated_lines = _estimate_total_bullet_lines(items, chars_per_line=chars_per_line)

    if bullet_count <= 1 and estimated_lines <= 2 and total_chars <= 90:
        state = "sparse"
    elif bullet_count <= 4 and estimated_lines <= 6 and total_chars <= 280:
        state = "normal"
    elif bullet_count <= 7 and estimated_lines <= 10 and total_chars <= 520:
        state = "dense"
    else:
        state = "overflow"

    return {
        "state": state,
        "bullet_count": bullet_count,
        "total_chars": total_chars,
        "estimated_lines": estimated_lines,
    }


def _capture_text_style(shape) -> dict:
    style = {
        "font_name": None,
        "font_size": None,
        "bold": None,
        "italic": None,
        "underline": None,
        "color_rgb": None,
        "alignment": None,
        "level": 0,
    }
    if not shape.has_text_frame or not shape.text_frame.paragraphs:
        return style

    paragraph = shape.text_frame.paragraphs[0]
    style["alignment"] = paragraph.alignment
    style["level"] = paragraph.level
    run = paragraph.runs[0] if paragraph.runs else None
    if run is None:
        return style

    font = run.font
    style["font_name"] = font.name
    style["font_size"] = font.size
    style["bold"] = font.bold
    style["italic"] = font.italic
    style["underline"] = font.underline
    if font.color is not None and font.color.rgb is not None:
        style["color_rgb"] = font.color.rgb
    return style


def _write_lines_with_format(
    shape,
    lines: list[str],
    *,
    style: dict,
    font_size_pt: float,
    line_spacing: float,
    space_after_pt: float,
) -> None:
    if not shape.has_text_frame:
        return

    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.clear()

    prepared_lines = [str(line).strip() for line in lines if str(line).strip()]
    if not prepared_lines:
        tf.text = ""
        return

    for line_index, line_value in enumerate(prepared_lines):
        paragraph = tf.paragraphs[0] if line_index == 0 else tf.add_paragraph()
        paragraph.alignment = style.get("alignment")
        paragraph.level = style.get("level", 0)
        paragraph.line_spacing = float(line_spacing)
        paragraph.space_after = Pt(space_after_pt)

        run = paragraph.add_run()
        run.text = line_value

        font = run.font
        font.name = style.get("font_name")
        font.bold = style.get("bold")
        font.italic = style.get("italic")
        font.underline = style.get("underline")
        if style.get("color_rgb") is not None:
            font.color.rgb = style.get("color_rgb")
        font.size = Pt(font_size_pt)


def _shape_bounds(shape) -> tuple[float, float, float, float]:
    left = _to_in(shape.left)
    top = _to_in(shape.top)
    width = _to_in(shape.width)
    height = _to_in(shape.height)
    return left, top, width, height


def _clamp(value: float, minimum: float, maximum: float) -> float:
    return max(minimum, min(maximum, value))


def _select_density_state(metrics: dict[str, int | str], area_height_in: float) -> str:
    state = str(metrics["state"])
    rules = _DENSITY_RULES.get(state, _DENSITY_RULES["overflow"])
    estimated_lines = int(metrics["estimated_lines"])
    font_size = float(rules["font_size_pt"])
    line_spacing = float(rules["line_spacing"])
    approximate_capacity = max(
        1,
        int((area_height_in * 72.0) / (font_size * line_spacing)),
    )
    if estimated_lines > approximate_capacity:
        return "overflow"
    return state


def _clone_shape(slide, source_shape):
    new_element = deepcopy(source_shape._element)
    slide.shapes._spTree.insert_element_before(new_element, "p:extLst")
    return slide.shapes[-1]


def _remove_shapes(helper, shapes_to_remove: list) -> None:
    for shape in shapes_to_remove:
        try:
            helper.remove_shape(shape)
        except Exception:
            continue


def _build_additional_text_shape(
    slide,
    prototype_shape,
    *,
    left_in: float,
    top_in: float,
    width_in: float,
    height_in: float,
):
    new_shape = slide.shapes.add_textbox(
        _to_emu(left_in),
        _to_emu(top_in),
        _to_emu(width_in),
        _to_emu(height_in),
    )
    return new_shape


def _collect_marker_shapes(
    *,
    helper,
    slide_index: int,
    bullet_left_emu: int,
    area_top_in: float,
    area_bottom_in: float,
) -> list:
    markers = []
    slide = helper.get_slide(slide_index)
    top_emu = _to_emu(area_top_in - 0.55)
    bottom_emu = _to_emu(area_bottom_in + 0.55)
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            continue
        if shape.width > _to_emu(1.0) or shape.height > _to_emu(1.0):
            continue
        if shape.left > bullet_left_emu:
            continue
        if shape.top < top_emu or shape.top > bottom_emu:
            continue
        markers.append(shape)
    markers.sort(key=lambda shape: shape.top)
    return markers


def _ensure_marker_shapes(
    *,
    helper,
    slide_index: int,
    marker_shapes: list,
    target_count: int,
) -> list:
    if target_count <= 0:
        _remove_shapes(helper, marker_shapes)
        return []

    working = list(marker_shapes)
    if len(working) > target_count:
        _remove_shapes(helper, working[target_count:])
        working = working[:target_count]
    elif len(working) < target_count and working:
        source = working[-1]
        slide = helper.get_slide(slide_index)
        while len(working) < target_count:
            try:
                clone_shape = _clone_shape(slide, source)
            except Exception:
                clone_shape = slide.shapes.add_shape(
                    MSO_AUTO_SHAPE_TYPE.OVAL,
                    source.left,
                    source.top,
                    source.width,
                    source.height,
                )
                clone_shape.fill.solid()
                clone_shape.fill.fore_color.rgb = RGBColor(59, 95, 191)
                clone_shape.line.fill.background()
            working.append(clone_shape)
    return working


def _scale_nearby_decoratives(
    *,
    helper,
    slide_index: int,
    protected_shape_ids: set[int],
    area_top_in: float,
    area_bottom_in: float,
    scale: float,
) -> None:
    if scale >= 0.999:
        return

    slide = helper.get_slide(slide_index)
    min_top_emu = _to_emu(area_top_in - 0.9)
    max_bottom_emu = _to_emu(area_bottom_in + 0.9)
    for shape in slide.shapes:
        if id(shape) in protected_shape_ids:
            continue
        if getattr(shape, "has_text_frame", False):
            continue
        # Keep large background/structure elements fixed to preserve template look.
        if shape.width >= _to_emu(9.0) or shape.height >= _to_emu(9.0):
            continue
        if shape.width <= _to_emu(1.0) or shape.height <= _to_emu(1.0):
            continue
        shape_bottom = shape.top + shape.height
        if shape_bottom < min_top_emu or shape.top > max_bottom_emu:
            continue
        center_x = shape.left + shape.width / 2
        center_y = shape.top + shape.height / 2
        new_width = max(_to_emu(0.8), int(round(shape.width * scale)))
        new_height = max(_to_emu(0.8), int(round(shape.height * scale)))
        try:
            shape.width = new_width
            shape.height = new_height
            shape.left = int(round(center_x - new_width / 2))
            shape.top = int(round(center_y - new_height / 2))
        except Exception:
            continue


def _scale_nearby_decoratives_by_name(
    *,
    helper,
    slide_index: int,
    protected_shape_names: set[str],
    area_top_in: float,
    area_bottom_in: float,
    scale: float,
) -> None:
    if scale >= 0.999:
        return

    slide = helper.get_slide(slide_index)
    min_top_emu = _to_emu(area_top_in - 0.9)
    max_bottom_emu = _to_emu(area_bottom_in + 0.9)
    for shape in slide.shapes:
        if str(getattr(shape, "name", "")) in protected_shape_names:
            continue
        if getattr(shape, "has_text_frame", False):
            continue
        if shape.width >= _to_emu(9.0) or shape.height >= _to_emu(9.0):
            continue
        if shape.width <= _to_emu(1.0) or shape.height <= _to_emu(1.0):
            continue
        shape_bottom = shape.top + shape.height
        if shape_bottom < min_top_emu or shape.top > max_bottom_emu:
            continue
        center_x = shape.left + shape.width / 2
        center_y = shape.top + shape.height / 2
        new_width = max(_to_emu(0.8), int(round(shape.width * scale)))
        new_height = max(_to_emu(0.8), int(round(shape.height * scale)))
        try:
            shape.width = new_width
            shape.height = new_height
            shape.left = int(round(center_x - new_width / 2))
            shape.top = int(round(center_y - new_height / 2))
        except Exception:
            continue


def _try_layout_solution(
    *,
    bullet_items: list[str],
    bullet_width_in: float,
    available_height_in: float,
    base_font_pt: float,
    min_font_pt: float,
    line_spacing: float,
    item_gap_in: float,
) -> dict | None:
    font_candidates = []
    start_font = int(round(base_font_pt))
    end_font = int(round(max(_MIN_READABLE_FONT_PT, min_font_pt)))
    for size in range(start_font, end_font - 1, -1):
        font_candidates.append(float(size))

    for font_size_pt in font_candidates:
        chars_per_line = _estimate_chars_per_line(bullet_width_in, font_size_pt)
        line_counts = [
            max(1, _estimate_wrapped_lines(item, chars_per_line=chars_per_line))
            for item in bullet_items
        ]
        line_height_in = _line_height_in(font_size_pt, line_spacing=line_spacing)
        item_heights = [
            max(0.24, lines * line_height_in + 0.02)
            for lines in line_counts
        ]
        required_height = sum(item_heights)
        if len(item_heights) > 1:
            required_height += item_gap_in * (len(item_heights) - 1)

        if required_height <= available_height_in + 1e-6:
            return {
                "font_size_pt": font_size_pt,
                "chars_per_line": chars_per_line,
                "line_counts": line_counts,
                "line_height_in": line_height_in,
                "item_heights": item_heights,
                "required_height_in": required_height,
            }

    return None


def _fits_in_shape(text: str, *, width_in: float, height_in: float, font_size_pt: float, line_spacing: float) -> tuple[bool, int, float]:
    chars_per_line = _estimate_chars_per_line(width_in, font_size_pt)
    lines = max(1, _estimate_wrapped_lines(text, chars_per_line=chars_per_line))
    required_height = lines * _line_height_in(font_size_pt, line_spacing=line_spacing) + 0.03
    return required_height <= height_in + 1e-6, lines, required_height


def _fit_font_for_shape(
    text: str,
    *,
    width_in: float,
    height_in: float,
    preferred_font_pt: float,
    min_font_pt: float,
    line_spacing: float,
) -> tuple[float, int] | None:
    start_font = int(round(preferred_font_pt))
    end_font = int(round(max(_MIN_READABLE_FONT_PT, min_font_pt)))
    for size in range(start_font, end_font - 1, -1):
        ok, lines, _ = _fits_in_shape(
            text,
            width_in=width_in,
            height_in=height_in,
            font_size_pt=float(size),
            line_spacing=line_spacing,
        )
        if ok:
            return float(size), lines
    return None


def _split_text_lines_for_layout(value: str) -> list[str]:
    normalized = re.sub(r"\r\n?", "\n", str(value or ""))
    lines = [line.strip() for line in normalized.split("\n") if line.strip()]
    return lines


def _estimate_wrapped_lines_multiline(value: str, *, chars_per_line: int) -> int:
    lines = _split_text_lines_for_layout(value)
    if not lines:
        compact = str(value or "").strip()
        if not compact:
            return 0
        return max(1, _estimate_wrapped_lines(compact, chars_per_line=chars_per_line))
    return sum(max(1, _estimate_wrapped_lines(line, chars_per_line=chars_per_line)) for line in lines)


def _should_replace_stage_candidate(
    current_font_pt: float | None,
    challenger_font_pt: float,
    *,
    min_gain_pt: float = 1.0,
) -> bool:
    if current_font_pt is None:
        return True
    gain = float(challenger_font_pt) - float(current_font_pt)
    return gain >= float(min_gain_pt)


def _is_story_with_image_and_points_template(shapes: dict) -> bool:
    return _STORY_IMG_FAMILY_REQUIRED_KEYS.issubset({str(key) for key in shapes.keys()})


def _template_key_from_shapes(shapes: dict) -> str:
    if not isinstance(shapes, dict):
        return ""
    value = shapes.get(_TEMPLATE_KEY_META_FIELD, "")
    return str(value).strip()


def _is_text_with_image_focus_repeat_template(shapes: dict) -> bool:
    if _template_key_from_shapes(shapes) != _IMG_FOCUS_REPEAT_TEMPLATE_KEY:
        return False
    return _IMG_FOCUS_REPEAT_REQUIRED_KEYS.issubset({str(key) for key in shapes.keys()})


def _apply_text_with_image_focus_repeat_layout(
    *,
    helper,
    slide_index: int,
    shapes: dict,
    slide_payload: dict,
) -> set[str]:
    if not isinstance(slide_payload, dict):
        return set()
    if not _is_text_with_image_focus_repeat_template(shapes):
        return set()

    required_fields = (
        "title_shape",
        "first_textbox",
        "second_textbox",
        "picture_placeholder",
    )
    resolved_shapes = {}
    for field in required_fields:
        shape_name = str(shapes[field]["name"])
        shape = helper.get_shape_by_name(slide_index, shape_name)
        if shape is None:
            return set()
        resolved_shapes[field] = shape

    title_shape = resolved_shapes["title_shape"]
    first_shape = resolved_shapes["first_textbox"]
    second_shape = resolved_shapes["second_textbox"]
    picture_shape = resolved_shapes["picture_placeholder"]
    flow_shapes = [resolved_shapes[field] for field in _IMG_FOCUS_REPEAT_FLOW_FIELDS]

    title_style = _capture_text_style(title_shape)
    flow_styles = {
        field: _capture_text_style(resolved_shapes[field])
        for field in _IMG_FOCUS_REPEAT_FLOW_FIELDS
    }

    title_text = str(slide_payload.get("title_shape", "") or "").strip()
    text_values = {
        field: str(slide_payload.get(field, "") or "").strip()
        for field in _IMG_FOCUS_REPEAT_FLOW_FIELDS
    }
    text_lines = {
        field: _split_text_lines_for_layout(text_values[field])
        for field in _IMG_FOCUS_REPEAT_FLOW_FIELDS
    }
    image_value = str(slide_payload.get("picture_placeholder", "") or "").strip()
    has_image = bool(image_value) and not _is_sentinel_media(image_value)

    slide_width_in = _to_in(helper.prs.slide_width)
    slide_height_in = _to_in(helper.prs.slide_height)

    # Keep image fixed in a reserved box so text never competes with it.
    fixed_image_box = (
        _to_in(picture_shape.left),
        _to_in(picture_shape.top),
        _to_in(picture_shape.width),
        _to_in(picture_shape.height),
    )
    if has_image:
        picture_shape.left = _to_emu(fixed_image_box[0])
        picture_shape.top = _to_emu(fixed_image_box[1])
        picture_shape.width = _to_emu(fixed_image_box[2])
        picture_shape.height = _to_emu(fixed_image_box[3])

    base_left_in = min(_to_in(shape.left) for shape in flow_shapes)
    base_right_in = max(_to_in(shape.left + shape.width) for shape in flow_shapes)
    base_top_in = min(_to_in(shape.top) for shape in flow_shapes)
    base_bottom_in = max(_to_in(shape.top + shape.height) for shape in flow_shapes)

    second_top_in = _to_in(second_shape.top)
    second_bottom_in = _to_in(second_shape.top + second_shape.height)
    first_top_in = _to_in(first_shape.top)
    base_gap_in = max(0.05, first_top_in - second_bottom_in)

    title_bottom_in = _to_in(title_shape.top + title_shape.height)
    image_left_in = fixed_image_box[0]
    image_top_in = fixed_image_box[1]
    image_bottom_in = fixed_image_box[1] + fixed_image_box[3]

    baseline_width_in = base_right_in - base_left_in
    estimate_width_in = baseline_width_in + (1.00 if not has_image else 0.22)
    estimate_chars = _estimate_chars_per_line(estimate_width_in, 20.0)
    estimated_total_lines = sum(
        _estimate_wrapped_lines_multiline(text_values[field], chars_per_line=estimate_chars)
        for field in _IMG_FOCUS_REPEAT_FLOW_FIELDS
        if text_values[field]
    )

    if estimated_total_lines <= 4:
        max_font_pt = 32.0
        preferred_font_pt = 30.0
    elif estimated_total_lines <= 8:
        max_font_pt = 30.0
        preferred_font_pt = 28.0
    elif estimated_total_lines <= 13:
        max_font_pt = 28.0
        preferred_font_pt = 26.0
    elif estimated_total_lines <= 18:
        max_font_pt = 26.0
        preferred_font_pt = 24.0
    else:
        max_font_pt = 24.0
        preferred_font_pt = 22.0
    min_font_pt = 18.0
    base_line_spacing = 1.00

    chosen = None
    chosen_font_pt: float | None = None
    for stage in _IMG_FOCUS_REPEAT_LAYOUT_STAGES:
        column_left_in = max(0.32, base_left_in - float(stage["column_shift_left_in"]))

        if has_image:
            right_limit_in = image_left_in - float(stage["image_text_gap_in"])
        else:
            right_limit_in = slide_width_in - 0.28
        column_right_in = min(
            right_limit_in,
            base_right_in + float(stage["column_expand_right_in"]),
        )
        if not has_image:
            column_right_in = min(
                slide_width_in - 0.22,
                column_right_in + float(stage["column_no_image_bonus_right_in"]),
            )
        column_width_in = column_right_in - column_left_in
        if column_width_in < 5.0:
            continue

        text_top_floor_in = title_bottom_in + 0.12
        column_top_in = max(
            text_top_floor_in,
            base_top_in - float(stage["top_raise_in"]),
        )
        if has_image:
            column_top_in = min(column_top_in, image_top_in - 0.05)
            column_top_in = max(text_top_floor_in, column_top_in)

        column_bottom_in = min(
            slide_height_in - 0.20,
            base_bottom_in + float(stage["bottom_extend_in"]),
        )
        if has_image:
            column_bottom_in = max(
                column_bottom_in,
                image_bottom_in + float(stage["below_image_extend_in"]),
            )
            column_bottom_in = min(slide_height_in - 0.18, column_bottom_in)

        available_height_in = column_bottom_in - column_top_in
        if available_height_in < 2.2:
            continue

        gap_in = max(
            float(stage["min_gap_in"]),
            base_gap_in * float(stage["gap_scale"]),
        )
        line_spacing = base_line_spacing * float(stage["line_spacing_factor"])

        stage_max_font = max(float(max_font_pt), float(preferred_font_pt))
        if bool(stage["allow_shrink"]):
            floor_font = float(min_font_pt)
        else:
            floor_font = float(preferred_font_pt)
        candidate_fonts = [
            float(size)
            for size in range(
                int(round(stage_max_font)),
                int(round(floor_font)) - 1,
                -1,
            )
        ]

        stage_candidate = None
        for candidate_font_pt in candidate_fonts:
            chars_per_line = _estimate_chars_per_line(column_width_in, candidate_font_pt)
            heights: list[float] = []
            line_counts: dict[str, int] = {}
            total_required = 0.0

            for field in _IMG_FOCUS_REPEAT_FLOW_FIELDS:
                value = text_values[field]
                if value:
                    lines = max(
                        1,
                        _estimate_wrapped_lines_multiline(value, chars_per_line=chars_per_line),
                    )
                    line_counts[field] = lines
                    required_h = lines * _line_height_in(candidate_font_pt, line_spacing=line_spacing) + 0.04
                    height_in = max(float(_IMG_FOCUS_REPEAT_MIN_HEIGHTS[field]), required_h)
                else:
                    line_counts[field] = 0
                    height_in = 0.16
                heights.append(height_in)
                total_required += height_in

            total_required += gap_in
            if total_required <= available_height_in + 1e-6:
                stage_candidate = {
                    "stage": stage,
                    "font_size_pt": candidate_font_pt,
                    "line_spacing": line_spacing,
                    "column_left_in": column_left_in,
                    "column_top_in": column_top_in,
                    "column_width_in": column_width_in,
                    "available_height_in": available_height_in,
                    "gap_in": gap_in,
                    "field_heights": heights,
                    "required_height_in": total_required,
                    "line_counts": line_counts,
                }
                break
        if stage_candidate is None:
            continue
        if _should_replace_stage_candidate(
            chosen_font_pt,
            float(stage_candidate["font_size_pt"]),
            min_gain_pt=1.0,
        ):
            chosen = stage_candidate
            chosen_font_pt = float(stage_candidate["font_size_pt"])

    if chosen is None:
        fail_width_in = max(5.0, baseline_width_in + (1.10 if not has_image else 0.36))
        estimated_lines_fail = sum(
            _estimate_wrapped_lines_multiline(
                text_values[field],
                chars_per_line=_estimate_chars_per_line(fail_width_in, min_font_pt),
            )
            for field in _IMG_FOCUS_REPEAT_FLOW_FIELDS
            if text_values[field]
        )
        raise ValueError(
            f"text_with_image_focus_repeat content does not fit this template safely on slide {slide_index + 1}. "
            f"estimated_lines={estimated_lines_fail}, min_readable_font={int(min_font_pt)}pt. "
            "Use slide splitting, template fallback, or upstream summarization."
        )

    stage = chosen["stage"]
    font_size_pt = float(chosen["font_size_pt"])
    line_spacing = float(chosen["line_spacing"])
    column_left_in = float(chosen["column_left_in"])
    column_top_in = float(chosen["column_top_in"])
    column_width_in = float(chosen["column_width_in"])
    available_height_in = float(chosen["available_height_in"])
    gap_in = float(chosen["gap_in"])
    field_heights = list(chosen["field_heights"])
    required_height_in = float(chosen["required_height_in"])
    line_counts = dict(chosen["line_counts"])

    protected_shape_names = {
        str(title_shape.name),
        str(first_shape.name),
        str(second_shape.name),
        str(picture_shape.name),
    }
    _scale_nearby_decoratives_by_name(
        helper=helper,
        slide_index=slide_index,
        protected_shape_names=protected_shape_names,
        area_top_in=column_top_in,
        area_bottom_in=column_top_in + available_height_in,
        scale=float(stage["decorative_scale"]),
    )

    total_lines = sum(line_counts.values())
    extra_space = max(0.0, available_height_in - required_height_in)
    if extra_space > 0:
        if total_lines <= 8:
            gap_ratio = 0.12
            gap_cap = 0.12
            height_share = 0.55
        elif total_lines <= 13:
            gap_ratio = 0.16
            gap_cap = 0.16
            height_share = 0.75
        else:
            gap_ratio = 0.20
            gap_cap = 0.20
            height_share = 0.92
        gap_extra = min(gap_cap, extra_space * gap_ratio)
        gap_in += gap_extra
        remaining = extra_space - gap_extra
        remaining_for_heights = max(0.0, remaining * height_share)
        if remaining_for_heights > 1e-6:
            second_weight = max(1.2, float(line_counts.get("second_textbox", 0)))
            first_weight = max(0.8, float(line_counts.get("first_textbox", 0)))
            total_weight = second_weight + first_weight
            field_heights[0] += remaining_for_heights * (second_weight / total_weight)
            field_heights[1] += remaining_for_heights * (first_weight / total_weight)

    used_height = field_heights[0] + field_heights[1] + gap_in
    residual_space = max(0.0, available_height_in - used_height)
    if total_lines <= 8:
        start_offset = min(0.14, residual_space * 0.30)
    elif total_lines <= 13:
        start_offset = min(0.10, residual_space * 0.20)
    else:
        start_offset = min(0.06, residual_space * 0.10)

    cursor_in = column_top_in + start_offset
    for idx, field in enumerate(_IMG_FOCUS_REPEAT_FLOW_FIELDS):
        shape = resolved_shapes[field]
        shape.left = _to_emu(column_left_in)
        shape.top = _to_emu(cursor_in)
        shape.width = _to_emu(column_width_in)
        shape.height = _to_emu(field_heights[idx])
        _write_lines_with_format(
            shape,
            text_lines[field] if text_lines[field] else [],
            style=flow_styles[field],
            font_size_pt=font_size_pt,
            line_spacing=line_spacing,
            space_after_pt=0.0,
        )
        cursor_in += field_heights[idx]
        if idx == 0:
            cursor_in += gap_in

    title_fit = _fit_font_for_shape(
        title_text,
        width_in=_to_in(title_shape.width),
        height_in=_to_in(title_shape.height),
        preferred_font_pt=max(32.0, font_size_pt + 2.0),
        min_font_pt=float(_MIN_TITLE_FONT_PT),
        line_spacing=1.0,
    )
    if title_fit is None:
        raise ValueError(
            f"text_with_image_focus_repeat title does not fit safely on slide {slide_index + 1}. "
            "Use title summarization or template fallback."
        )
    title_font_pt, _ = title_fit
    _write_lines_with_format(
        title_shape,
        [title_text] if title_text else [],
        style=title_style,
        font_size_pt=float(title_font_pt),
        line_spacing=1.0,
        space_after_pt=0.0,
    )

    return {
        "title_shape",
        "first_textbox",
        "second_textbox",
    }


def _apply_story_with_image_and_points_layout(
    *,
    helper,
    slide_index: int,
    shapes: dict,
    slide_payload: dict,
) -> set[str]:
    if not isinstance(slide_payload, dict):
        return set()
    if not _is_story_with_image_and_points_template(shapes):
        return set()

    required_fields = (
        "title_shape",
        "first_textbox",
        "second_textbox",
        "third_textbox",
        "fourth_textbox",
        "fifth_textbox",
        "picture_placeholder",
    )
    resolved_shapes = {}
    for field in required_fields:
        shape_name = str(shapes[field]["name"])
        shape = helper.get_shape_by_name(slide_index, shape_name)
        if shape is None:
            return set()
        resolved_shapes[field] = shape

    title_shape = resolved_shapes["title_shape"]
    picture_shape = resolved_shapes["picture_placeholder"]
    flow_shapes = [resolved_shapes[field] for field in _STORY_IMG_FLOW_FIELDS]
    flow_styles = {field: _capture_text_style(resolved_shapes[field]) for field in _STORY_IMG_FLOW_FIELDS}
    title_style = _capture_text_style(title_shape)

    text_values = {field: str(slide_payload.get(field, "") or "").strip() for field in _STORY_IMG_FLOW_FIELDS}
    title_text = str(slide_payload.get("title_shape", "") or "").strip()
    image_value = str(slide_payload.get("picture_placeholder", "") or "").strip()
    has_image = bool(image_value) and not _is_sentinel_media(image_value)

    slide_width_in = _to_in(helper.prs.slide_width)
    slide_height_in = _to_in(helper.prs.slide_height)

    # Reserve a fixed image region for this family whenever an image is present.
    fixed_image_box = (
        _to_in(picture_shape.left),
        _to_in(picture_shape.top),
        _to_in(picture_shape.width),
        _to_in(picture_shape.height),
    )
    if has_image:
        picture_shape.left = _to_emu(fixed_image_box[0])
        picture_shape.top = _to_emu(fixed_image_box[1])
        picture_shape.width = _to_emu(fixed_image_box[2])
        picture_shape.height = _to_emu(fixed_image_box[3])

    base_left_in = min(_to_in(shape.left) for shape in flow_shapes)
    base_right_in = max(_to_in(shape.left + shape.width) for shape in flow_shapes)
    base_top_in = min(_to_in(shape.top) for shape in flow_shapes)
    base_bottom_in = max(_to_in(shape.top + shape.height) for shape in flow_shapes)

    base_positions = [(_to_in(resolved_shapes[field].top), _to_in(resolved_shapes[field].height)) for field in _STORY_IMG_FLOW_FIELDS]
    base_gaps: list[float] = []
    for index in range(len(base_positions) - 1):
        current_top, current_height = base_positions[index]
        next_top, _ = base_positions[index + 1]
        base_gaps.append(max(0.04, next_top - (current_top + current_height)))

    title_right_in = _to_in(title_shape.left + title_shape.width)
    image_right_in = fixed_image_box[0] + fixed_image_box[2]

    baseline_width = base_right_in - base_left_in
    estimate_width = baseline_width + (1.2 if not has_image else 0.2)
    estimate_chars = _estimate_chars_per_line(estimate_width, 20.0)
    estimated_total_lines = sum(
        _estimate_wrapped_lines(text_values[field], chars_per_line=estimate_chars)
        for field in _STORY_IMG_FLOW_FIELDS
        if text_values[field]
    )

    if estimated_total_lines <= 6:
        max_font_pt = 32.0
        preferred_font_pt = 30.0
    elif estimated_total_lines <= 10:
        max_font_pt = 30.0
        preferred_font_pt = 28.0
    elif estimated_total_lines <= 15:
        max_font_pt = 28.0
        preferred_font_pt = 26.0
    else:
        max_font_pt = 26.0
        preferred_font_pt = 24.0
    min_font_pt = 19.0
    base_line_spacing = 1.02

    chosen = None
    chosen_font_pt: float | None = None
    for stage in _STORY_IMG_LAYOUT_STAGES:
        left_shift_in = (
            float(stage["column_shift_left_with_image_in"])
            if has_image
            else float(stage["column_shift_left_without_image_in"])
        )
        column_left_in = base_left_in - left_shift_in
        min_left_in = title_right_in + 0.35
        if has_image:
            min_left_in = max(min_left_in, image_right_in + 0.45)
        column_left_in = max(min_left_in, column_left_in)

        column_right_in = min(slide_width_in - 0.35, base_right_in + float(stage["column_expand_right_in"]))
        if not has_image:
            column_right_in = min(slide_width_in - 0.25, column_right_in + 0.30)
        column_width_in = column_right_in - column_left_in
        if column_width_in < 7.4:
            continue

        column_top_in = max(0.72, base_top_in - float(stage["top_raise_in"]))
        column_bottom_in = min(slide_height_in - 0.28, base_bottom_in + float(stage["bottom_extend_in"]))
        if not has_image:
            column_bottom_in = min(slide_height_in - 0.20, column_bottom_in + 0.18)
        available_height_in = column_bottom_in - column_top_in
        if available_height_in < 2.7:
            continue

        gaps = [max(float(stage["min_gap_in"]), gap * float(stage["gap_scale"])) for gap in base_gaps]
        line_spacing = base_line_spacing * float(stage["line_spacing_factor"])

        stage_max_font = max(float(max_font_pt), float(preferred_font_pt))
        if stage is _STORY_IMG_LAYOUT_STAGES[-1]:
            floor_font = float(min_font_pt)
        else:
            floor_font = float(preferred_font_pt)
        candidate_fonts = [
            float(size)
            for size in range(
                int(round(stage_max_font)),
                int(round(floor_font)) - 1,
                -1,
            )
        ]

        stage_candidate = None
        for candidate_font_pt in candidate_fonts:
            chars_per_line = _estimate_chars_per_line(column_width_in, candidate_font_pt)
            field_heights: list[float] = []
            line_counts: dict[str, int] = {}
            total_required = 0.0

            for field in _STORY_IMG_FLOW_FIELDS:
                value = text_values[field]
                if value:
                    lines = max(1, _estimate_wrapped_lines(value, chars_per_line=chars_per_line))
                    line_counts[field] = lines
                    required_h = lines * _line_height_in(candidate_font_pt, line_spacing=line_spacing) + 0.04
                    height_in = max(float(_STORY_IMG_MIN_HEIGHTS[field]), required_h)
                else:
                    line_counts[field] = 0
                    height_in = 0.16
                field_heights.append(height_in)
                total_required += height_in

            total_required += sum(gaps)
            if total_required <= available_height_in + 1e-6:
                stage_candidate = {
                    "stage": stage,
                    "font_size_pt": candidate_font_pt,
                    "line_spacing": line_spacing,
                    "column_left_in": column_left_in,
                    "column_top_in": column_top_in,
                    "column_width_in": column_width_in,
                    "available_height_in": available_height_in,
                    "field_heights": field_heights,
                    "gaps": gaps,
                    "required_height_in": total_required,
                    "line_counts": line_counts,
                }
                break
        if stage_candidate is None:
            continue
        if _should_replace_stage_candidate(
            chosen_font_pt,
            float(stage_candidate["font_size_pt"]),
            min_gain_pt=1.0,
        ):
            chosen = stage_candidate
            chosen_font_pt = float(stage_candidate["font_size_pt"])

    if chosen is None:
        estimate_width_fail = max(7.4, baseline_width + (1.4 if not has_image else 0.4))
        estimated_lines_fail = sum(
            _estimate_wrapped_lines(text_values[field], chars_per_line=_estimate_chars_per_line(estimate_width_fail, min_font_pt))
            for field in _STORY_IMG_FLOW_FIELDS
            if text_values[field]
        )
        raise ValueError(
            f"Story-with-image-and-points content does not fit this template safely on slide {slide_index + 1}. "
            f"estimated_lines={estimated_lines_fail}, min_readable_font={int(min_font_pt)}pt. "
            "Use slide splitting, template fallback, or upstream summarization."
        )

    stage = chosen["stage"]
    font_size_pt = float(chosen["font_size_pt"])
    line_spacing = float(chosen["line_spacing"])
    column_left_in = float(chosen["column_left_in"])
    column_top_in = float(chosen["column_top_in"])
    column_width_in = float(chosen["column_width_in"])
    field_heights = list(chosen["field_heights"])
    gaps = list(chosen["gaps"])
    available_height_in = float(chosen["available_height_in"])
    required_height_in = float(chosen["required_height_in"])

    protected_shape_ids = {id(title_shape), id(picture_shape), *(id(shape) for shape in flow_shapes)}
    _scale_nearby_decoratives(
        helper=helper,
        slide_index=slide_index,
        protected_shape_ids=protected_shape_ids,
        area_top_in=column_top_in,
        area_bottom_in=column_top_in + available_height_in,
        scale=float(stage["decorative_scale"]),
    )

    extra_space = max(0.0, available_height_in - required_height_in)
    total_lines = sum(chosen["line_counts"].values())
    if extra_space > 0:
        if total_lines <= 8:
            gap_share = 0.25
            per_gap_cap = 0.14
        elif total_lines <= 13:
            gap_share = 0.32
            per_gap_cap = 0.16
        else:
            gap_share = 0.40
            per_gap_cap = 0.20
        extra_for_gaps = extra_space * gap_share if gaps else 0.0
        consumed_gap_extra = 0.0
        if gaps:
            per_gap = extra_for_gaps / len(gaps)
            for idx in range(len(gaps)):
                add_gap = min(per_gap_cap, per_gap)
                gaps[idx] += add_gap
                consumed_gap_extra += add_gap

        remaining_extra = extra_space - consumed_gap_extra
        if remaining_extra > 1e-6:
            weights = []
            for field in _STORY_IMG_FLOW_FIELDS:
                weights.append(max(0.5, float(chosen["line_counts"].get(field, 0))))
            total_weight = sum(weights)
            if total_weight > 0:
                for idx, weight in enumerate(weights):
                    field_heights[idx] += remaining_extra * (weight / total_weight)

    used_height = sum(field_heights) + sum(gaps)
    residual_space = max(0.0, available_height_in - used_height)
    if total_lines <= 8:
        start_offset = min(0.24, residual_space * 0.34)
    elif total_lines <= 13:
        start_offset = min(0.15, residual_space * 0.22)
    else:
        start_offset = min(0.06, residual_space * 0.10)

    cursor_in = column_top_in + start_offset
    for index, field in enumerate(_STORY_IMG_FLOW_FIELDS):
        shape = resolved_shapes[field]
        shape.left = _to_emu(column_left_in)
        shape.top = _to_emu(cursor_in)
        shape.width = _to_emu(column_width_in)
        shape.height = _to_emu(field_heights[index])
        _write_lines_with_format(
            shape,
            [text_values[field]] if text_values[field] else [],
            style=flow_styles[field],
            font_size_pt=font_size_pt,
            line_spacing=line_spacing,
            space_after_pt=0.5,
        )
        gap = gaps[index] if index < len(gaps) else 0.0
        cursor_in += field_heights[index] + gap

    title_fit = _fit_font_for_shape(
        title_text,
        width_in=_to_in(title_shape.width),
        height_in=_to_in(title_shape.height),
        preferred_font_pt=max(32.0, font_size_pt + 2.0),
        min_font_pt=float(_MIN_TITLE_FONT_PT),
        line_spacing=1.0,
    )
    if title_fit is None:
        raise ValueError(
            f"Story-with-image-and-points title does not fit safely on slide {slide_index + 1}. "
            "Use title summarization or template fallback."
        )
    title_font_pt, _ = title_fit
    _write_lines_with_format(
        title_shape,
        [title_text] if title_text else [],
        style=title_style,
        font_size_pt=float(title_font_pt),
        line_spacing=1.0,
        space_after_pt=0.0,
    )

    return {"title_shape", *set(_STORY_IMG_FLOW_FIELDS)}


def _apply_adaptive_bullet_layout(
    *,
    helper,
    slide_index: int,
    shapes: dict,
    slide_payload: dict,
) -> set[str]:
    if not isinstance(slide_payload, dict):
        return set()

    bullet_shape_entries = [
        {
            "shape_key": str(shape_key),
            "shape_name": str(shape_info.get("name")),
            "shape": helper.get_shape_by_name(slide_index, str(shape_info.get("name"))),
        }
        for shape_key, shape_info in shapes.items()
        if "bullet" in str(shape_key).lower()
    ]
    bullet_shape_entries = [entry for entry in bullet_shape_entries if entry["shape"] is not None]
    if not bullet_shape_entries:
        return set()

    bullet_shape_entries.sort(key=lambda entry: entry["shape"].top)
    bullet_shapes = [entry["shape"] for entry in bullet_shape_entries]
    bullet_keys = _ordered_bullet_keys(slide_payload)
    bullet_items = _collect_bullet_items(slide_payload, bullet_keys)
    if not bullet_items:
        handled_keys = {entry["shape_key"] for entry in bullet_shape_entries}
        for key in bullet_keys:
            if key in slide_payload:
                slide_payload[key] = ""
        return handled_keys

    first_text_shape = None
    second_text_shape = None
    title_shape = None
    if "first_textbox" in shapes:
        first_text_shape = helper.get_shape_by_name(slide_index, shapes["first_textbox"]["name"])
    if "second_textbox" in shapes:
        second_text_shape = helper.get_shape_by_name(slide_index, shapes["second_textbox"]["name"])
    if "title_shape" in shapes:
        title_shape = helper.get_shape_by_name(slide_index, shapes["title_shape"]["name"])
    if second_text_shape is None:
        return set()

    first_text_style = _capture_text_style(first_text_shape) if first_text_shape is not None else None
    second_text_style = _capture_text_style(second_text_shape)
    first_text_value = str(slide_payload.get("first_textbox", "") or "").strip()
    second_text_value = str(slide_payload.get("second_textbox", "") or "").strip()

    base_left_in = min(_to_in(shape.left) for shape in bullet_shapes)
    base_width_in = max(_to_in(shape.width) for shape in bullet_shapes)
    default_top_in = min(_to_in(shape.top) for shape in bullet_shapes)
    default_bottom_in = max(_to_in(shape.top + shape.height) for shape in bullet_shapes)
    default_area_height_in = max(0.8, default_bottom_in - default_top_in)
    slide_width_in = _to_in(helper.prs.slide_width)
    slide_height_in = _to_in(helper.prs.slide_height)

    if first_text_shape is not None and first_text_value:
        first_current_height_in = _to_in(first_text_shape.height)
        first_width_in = _to_in(first_text_shape.width)
        first_len = len(first_text_value)
        if first_len <= 95:
            preview_font_pt = 26.0
        elif first_len <= 180:
            preview_font_pt = 23.0
        else:
            preview_font_pt = 22.0
        preview_lines = max(
            1,
            _estimate_wrapped_lines(
                first_text_value,
                chars_per_line=_estimate_chars_per_line(first_width_in, preview_font_pt),
            ),
        )
        preview_required_h = preview_lines * _line_height_in(preview_font_pt, line_spacing=1.0) + 0.22
        if first_len <= 140:
            target_first_h = _clamp(preview_required_h, 1.10, first_current_height_in)
            first_text_shape.height = _to_emu(target_first_h)

    first_bottom_in = _to_in(first_text_shape.top + first_text_shape.height) if first_text_shape is not None else (default_top_in - 0.12)
    second_default_top_in = _to_in(second_text_shape.top)
    second_default_height_in = _to_in(second_text_shape.height)

    base_chars = _estimate_chars_per_line(base_width_in, _BASE_FIT_FONT_PT)
    density_metrics = _classify_bullet_density(bullet_items, chars_per_line=base_chars)
    density_state = _select_density_state(density_metrics, default_area_height_in)
    density_rules = _DENSITY_RULES.get(density_state, _DENSITY_RULES["overflow"])
    preferred_font_pt = float(density_rules["font_size_pt"])
    if density_state == "sparse":
        max_font_pt = 32.0
    elif density_state == "normal":
        max_font_pt = 30.0
    elif density_state == "dense":
        max_font_pt = 27.0
    else:
        max_font_pt = 24.0
    minimum_font_pt = float(max(_MIN_READABLE_FONT_PT, density_rules["min_font_pt"]))
    text_style = _capture_text_style(bullet_shapes[0])
    title_text = str(slide_payload.get("title_shape", "") or "").strip()
    title_style = _capture_text_style(title_shape) if title_shape is not None else None

    chosen = None
    chosen_font_pt: float | None = None
    for stage in _LAYOUT_STAGES:
        second_min_height_in = float(stage["second_min_height_in"])
        second_height_reduce_in = float(stage["second_height_reduce_in"])
        if density_state == "sparse":
            second_min_height_in = max(0.98, second_min_height_in - 0.28)
            second_height_reduce_in += 0.18
        second_height_in = max(
            second_min_height_in,
            second_default_height_in - second_height_reduce_in,
        )
        max_second_top_in = min(
            second_default_top_in + float(stage["second_push_down_in"]),
            slide_height_in - second_height_in - 0.10,
        )
        area_top_in = max(first_bottom_in + 0.05, default_top_in - float(stage["top_raise_in"]))
        area_bottom_in = max(area_top_in + 0.6, max_second_top_in - float(stage["below_gap_in"]))
        available_height_in = area_bottom_in - area_top_in
        bullet_width_in = _clamp(
            base_width_in + float(stage["width_expand_in"]),
            base_width_in,
            slide_width_in - base_left_in - 0.35,
        )

        if stage is _LAYOUT_STAGES[-1]:
            floor_font = minimum_font_pt
        else:
            floor_font = preferred_font_pt
        fit_solution = _try_layout_solution(
            bullet_items=bullet_items,
            bullet_width_in=bullet_width_in,
            available_height_in=available_height_in,
            base_font_pt=max(max_font_pt, preferred_font_pt),
            min_font_pt=floor_font,
            line_spacing=float(density_rules["line_spacing"]) * float(stage["line_spacing_factor"]),
            item_gap_in=float(stage["item_gap_in"]),
        )
        if fit_solution is None:
            continue

        stage_candidate = {
            "stage": stage,
            "area_top_in": area_top_in,
            "area_bottom_in": area_bottom_in,
            "available_height_in": available_height_in,
            "second_top_in": max_second_top_in,
            "second_height_in": second_height_in,
            "bullet_width_in": bullet_width_in,
            "fit": fit_solution,
        }
        if _should_replace_stage_candidate(
            chosen_font_pt,
            float(fit_solution["font_size_pt"]),
            min_gain_pt=1.0,
        ):
            chosen = stage_candidate
            chosen_font_pt = float(fit_solution["font_size_pt"])

    if chosen is None:
        estimated_lines = _estimate_total_bullet_lines(
            bullet_items,
            chars_per_line=_estimate_chars_per_line(base_width_in + 0.24, _MIN_READABLE_FONT_PT),
        )
        raise ValueError(
            f"Bullet content does not fit this template safely on slide {slide_index + 1}. "
            f"bullets={len(bullet_items)}, estimated_lines={estimated_lines}, "
            f"min_readable_font={_MIN_READABLE_FONT_PT}pt. "
            "Use slide splitting, template fallback, or upstream content reduction."
        )

    stage = chosen["stage"]
    fit = chosen["fit"]
    area_top_in = float(chosen["area_top_in"])
    available_height_in = float(chosen["available_height_in"])
    bullet_width_in = float(chosen["bullet_width_in"])
    line_spacing = float(density_rules["line_spacing"]) * float(stage["line_spacing_factor"])
    item_gap_in = float(stage["item_gap_in"])
    font_size_pt = float(fit["font_size_pt"])
    item_heights = list(fit["item_heights"])
    line_height_in = float(fit["line_height_in"])

    second_text_shape.top = _to_emu(float(chosen["second_top_in"]))
    second_text_shape.height = _to_emu(float(chosen["second_height_in"]))

    bullet_count = len(bullet_items)
    base_text_shapes = list(bullet_shapes)
    if len(base_text_shapes) > bullet_count:
        _remove_shapes(helper, base_text_shapes[bullet_count:])
        base_text_shapes = base_text_shapes[:bullet_count]

    slide = helper.get_slide(slide_index)
    while len(base_text_shapes) < bullet_count:
        base_text_shapes.append(
            _build_additional_text_shape(
                slide,
                prototype_shape=bullet_shapes[-1],
                left_in=base_left_in,
                top_in=area_top_in,
                width_in=bullet_width_in,
                height_in=max(0.24, line_height_in),
            )
        )

    marker_shapes = _collect_marker_shapes(
        helper=helper,
        slide_index=slide_index,
        bullet_left_emu=min(shape.left for shape in base_text_shapes),
        area_top_in=default_top_in,
        area_bottom_in=default_bottom_in,
    )
    marker_shapes = _ensure_marker_shapes(
        helper=helper,
        slide_index=slide_index,
        marker_shapes=marker_shapes,
        target_count=bullet_count,
    )

    protected_shape_ids = {id(shape) for shape in base_text_shapes}
    protected_shape_ids.update(id(shape) for shape in marker_shapes)
    protected_shape_ids.add(id(second_text_shape))
    _scale_nearby_decoratives(
        helper=helper,
        slide_index=slide_index,
        protected_shape_ids=protected_shape_ids,
        area_top_in=area_top_in,
        area_bottom_in=area_top_in + available_height_in,
        scale=float(stage["decorative_scale"]),
    )

    required_height_in = float(fit["required_height_in"])
    remaining_space = max(0.0, available_height_in - required_height_in)
    if remaining_space > 1e-6 and bullet_count > 0:
        if density_state == "sparse":
            gap_share = 0.30
            per_gap_cap = 0.18
        elif density_state == "normal":
            gap_share = 0.24
            per_gap_cap = 0.14
        else:
            gap_share = 0.16
            per_gap_cap = 0.10
        gap_extra_total = 0.0
        if bullet_count > 1:
            per_gap_add = min(per_gap_cap, (remaining_space * gap_share) / (bullet_count - 1))
            gap_extra_total = per_gap_add * (bullet_count - 1)
            item_gap_in += per_gap_add

        remaining_height_extra = max(0.0, remaining_space - gap_extra_total)
        if remaining_height_extra > 1e-6:
            line_weights = [max(1.0, float(lines)) for lines in fit["line_counts"]]
            weight_sum = sum(line_weights)
            if weight_sum > 0:
                for idx, weight in enumerate(line_weights):
                    item_heights[idx] += remaining_height_extra * (weight / weight_sum)

    bullet_stack_height_in = sum(item_heights)
    if bullet_count > 1:
        bullet_stack_height_in += item_gap_in * (bullet_count - 1)
    residual_after_stack_in = max(0.0, available_height_in - bullet_stack_height_in)

    if density_state == "sparse":
        start_top_in = area_top_in + min(0.22, residual_after_stack_in * 0.35)
    elif density_state == "normal":
        start_top_in = area_top_in + min(0.14, residual_after_stack_in * 0.18)
    else:
        start_top_in = area_top_in + min(0.07, residual_after_stack_in * 0.10)

    top_cursor_in = start_top_in
    for shape, item_text, item_height_in in zip(base_text_shapes, bullet_items, item_heights):
        shape.left = _to_emu(base_left_in)
        shape.top = _to_emu(top_cursor_in)
        shape.width = _to_emu(bullet_width_in)
        shape.height = _to_emu(item_height_in)
        _write_lines_with_format(
            shape,
            [item_text],
            style=text_style,
            font_size_pt=font_size_pt,
            line_spacing=line_spacing,
            space_after_pt=0.0,
        )
        top_cursor_in += item_height_in + item_gap_in

    if marker_shapes:
        marker_base_left = min(_to_in(shape.left) for shape in marker_shapes)
        marker_base_width = sum(_to_in(shape.width) for shape in marker_shapes) / len(marker_shapes)
        marker_base_height = sum(_to_in(shape.height) for shape in marker_shapes) / len(marker_shapes)
        marker_scale = float(stage["marker_scale"]) * _clamp(font_size_pt / _BASE_FIT_FONT_PT, 0.82, 1.02)
        marker_width_in = max(0.14, marker_base_width * marker_scale)
        marker_height_in = max(0.14, marker_base_height * marker_scale)

        marker_cursor_in = start_top_in
        for marker_shape, item_height_in in zip(marker_shapes, item_heights):
            marker_anchor_top = marker_cursor_in + min(item_height_in, line_height_in) * 0.52
            marker_shape.left = _to_emu(marker_base_left)
            marker_shape.width = _to_emu(marker_width_in)
            marker_shape.height = _to_emu(marker_height_in)
            marker_shape.top = _to_emu(marker_anchor_top - marker_height_in / 2)
            marker_cursor_in += item_height_in + item_gap_in

    if first_text_shape is not None:
        first_text_len = len(first_text_value)
        if first_text_len <= 95:
            first_preferred_pt = 30.0
        elif first_text_len <= 180:
            first_preferred_pt = 28.0
        else:
            first_preferred_pt = 26.0
        first_fit = _fit_font_for_shape(
            first_text_value,
            width_in=_to_in(first_text_shape.width),
            height_in=_to_in(first_text_shape.height),
            preferred_font_pt=first_preferred_pt,
            min_font_pt=18.0,
            line_spacing=1.0,
        )
        if first_fit is None:
            raise ValueError(
                f"Bullet slide first_textbox does not fit safely on slide {slide_index + 1}. "
                "Use text summarization or template fallback."
            )
        first_font_pt, _ = first_fit
        _write_lines_with_format(
            first_text_shape,
            [first_text_value] if first_text_value else [],
            style=first_text_style or {},
            font_size_pt=float(first_font_pt),
            line_spacing=1.0,
            space_after_pt=0.0,
        )

    second_text_len = len(second_text_value)
    if second_text_len <= 120:
        second_preferred_pt = 29.0
    elif second_text_len <= 210:
        second_preferred_pt = 27.0
    else:
        second_preferred_pt = 25.0
    second_fit = _fit_font_for_shape(
        second_text_value,
        width_in=_to_in(second_text_shape.width),
        height_in=_to_in(second_text_shape.height),
        preferred_font_pt=second_preferred_pt,
        min_font_pt=18.0,
        line_spacing=1.0,
    )
    if second_fit is None:
        raise ValueError(
            f"Bullet slide second_textbox does not fit safely on slide {slide_index + 1}. "
            "Use text summarization or template fallback."
        )
    second_font_pt, _ = second_fit
    _write_lines_with_format(
        second_text_shape,
        [second_text_value] if second_text_value else [],
        style=second_text_style or {},
        font_size_pt=float(second_font_pt),
        line_spacing=1.0,
        space_after_pt=0.0,
    )

    if title_shape is not None:
        title_fit = _fit_font_for_shape(
            title_text,
            width_in=_to_in(title_shape.width),
            height_in=_to_in(title_shape.height),
            preferred_font_pt=max(32.0, float(_MIN_TITLE_FONT_PT)),
            min_font_pt=float(_MIN_TITLE_FONT_PT),
            line_spacing=1.0,
        )
        if title_fit is None:
            raise ValueError(
                f"Bullet slide title does not fit safely on slide {slide_index + 1}. "
                "Use title summarization or template fallback."
            )
        title_font_pt, _ = title_fit
        _write_lines_with_format(
            title_shape,
            [title_text] if title_text else [],
            style=title_style or {},
            font_size_pt=float(title_font_pt),
            line_spacing=1.0,
            space_after_pt=0.0,
        )

    handled_keys = {entry["shape_key"] for entry in bullet_shape_entries}
    if title_shape is not None and "title_shape" in shapes:
        handled_keys.add("title_shape")
    if first_text_shape is not None and "first_textbox" in shapes:
        handled_keys.add("first_textbox")
    if "second_textbox" in shapes:
        handled_keys.add("second_textbox")
    for key in bullet_keys:
        if key in slide_payload:
            slide_payload[key] = ""
    return handled_keys


def _is_sentinel_media(value: str) -> bool:
    normalized = str(value).strip().lower()
    if not normalized:
        return False

    if normalized == "manim_vid":
        return True

    if normalized.startswith("[") and normalized.endswith("]"):
        normalized = normalized[1:-1].strip()

    if normalized.startswith("video:"):
        media_value = normalized.split(":", 1)[1].strip()
        return media_value == "manim_vid"

    return False


def _normalize_media_reference(image_name: str) -> str:
    normalized = str(image_name).strip()
    if not normalized:
        return normalized

    if normalized.startswith("[") and normalized.endswith("]"):
        normalized = normalized[1:-1].strip()

    if normalized.lower().startswith("img:"):
        normalized = normalized.split(":", 1)[1].strip()

    return normalized


def fill_image_placeholder(helper, slide_index: int, placeholder_name: str, image_path: str):
    if not image_path:
        return

    shape = helper.get_shape_by_name(slide_index, placeholder_name)
    if not shape:
        print(f"Placeholder not found: {placeholder_name} on slide {slide_index + 1}")
        return

    if not Path(image_path).is_file():
        print(f"Image file not found: {image_path}")
        return

    left = shape.left
    top = shape.top
    width = shape.width
    height = shape.height

    helper.remove_shape(shape)

    slide = helper.get_slide(slide_index)
    slide.shapes.add_picture(image_path, left, top, width, height)


def _prepare_image_for_cover_box(
    image_path: str,
    *,
    target_width_emu: int,
    target_height_emu: int,
) -> str:
    """Create a centered cover-cropped image matching the target box aspect ratio."""
    source_path = Path(image_path)
    if not source_path.is_file():
        return image_path

    target_width_in = _to_in(target_width_emu)
    target_height_in = _to_in(target_height_emu)
    if target_width_in <= 0 or target_height_in <= 0:
        return image_path

    target_width_px = max(80, int(round(target_width_in * 220)))
    target_height_px = max(80, int(round(target_height_in * 220)))

    try:
        from PIL import Image
    except Exception as exc:
        raise RuntimeError(
            "Pillow is required for story_with_image_and_points image preprocessing."
        ) from exc

    with Image.open(source_path) as image:
        source_width, source_height = image.size
        if source_width <= 0 or source_height <= 0:
            return image_path

        cover_scale = max(
            float(target_width_px) / float(source_width),
            float(target_height_px) / float(source_height),
        )
        resized_width = max(target_width_px, int(math.ceil(source_width * cover_scale)))
        resized_height = max(target_height_px, int(math.ceil(source_height * cover_scale)))
        resized = image.resize((resized_width, resized_height), Image.Resampling.LANCZOS)

        crop_left = max(0, int(round((resized_width - target_width_px) / 2.0)))
        crop_top = max(0, int(round((resized_height - target_height_px) / 2.0)))
        crop_box = (
            crop_left,
            crop_top,
            crop_left + target_width_px,
            crop_top + target_height_px,
        )
        prepared = resized.crop(crop_box)
        if prepared.mode not in ("RGB", "RGBA"):
            prepared = prepared.convert("RGB")

    prepared_name = (
        f"{source_path.stem}__story_cover_{target_width_px}x{target_height_px}.png"
    )
    prepared_path = source_path.with_name(prepared_name)
    prepared.save(prepared_path, format="PNG", optimize=True)
    return str(prepared_path)


def _prepare_image_for_contain_box(
    image_path: str,
    *,
    target_width_emu: int,
    target_height_emu: int,
) -> str:
    """Create a centered fit-inside image with transparent margins (no crop, no distortion)."""
    source_path = Path(image_path)
    if not source_path.is_file():
        return image_path

    target_width_in = _to_in(target_width_emu)
    target_height_in = _to_in(target_height_emu)
    if target_width_in <= 0 or target_height_in <= 0:
        return image_path

    target_width_px = max(80, int(round(target_width_in * 220)))
    target_height_px = max(80, int(round(target_height_in * 220)))

    try:
        from PIL import Image
    except Exception as exc:
        raise RuntimeError(
            "Pillow is required for text_with_image_focus_repeat image preprocessing."
        ) from exc

    with Image.open(source_path) as image:
        source_width, source_height = image.size
        if source_width <= 0 or source_height <= 0:
            return image_path

        fit_scale = min(
            float(target_width_px) / float(source_width),
            float(target_height_px) / float(source_height),
        )
        resized_width = max(1, int(round(source_width * fit_scale)))
        resized_height = max(1, int(round(source_height * fit_scale)))

        resized = image.resize((resized_width, resized_height), Image.Resampling.LANCZOS)
        if resized.mode != "RGBA":
            resized = resized.convert("RGBA")

        canvas = Image.new("RGBA", (target_width_px, target_height_px), (0, 0, 0, 0))
        paste_left = max(0, int(round((target_width_px - resized_width) / 2.0)))
        paste_top = max(0, int(round((target_height_px - resized_height) / 2.0)))
        canvas.paste(resized, (paste_left, paste_top), resized)

    prepared_name = f"{source_path.stem}__focus_repeat_fit_{target_width_px}x{target_height_px}.png"
    prepared_path = source_path.with_name(prepared_name)
    canvas.save(prepared_path, format="PNG", optimize=True)
    return str(prepared_path)


def _candidate_media_json_paths(project_id, image_name: str, settings=None) -> list[Path]:
    resolved_settings = settings or config
    paths = ProjectPaths(project_id=project_id, settings=resolved_settings)
    media_folder = paths.dir("media", prefer_existing=True)
    normalized_name = _normalize_media_reference(image_name)
    if not normalized_name:
        return []

    raw_path = Path(normalized_name)
    candidates: list[Path] = []
    seen: set[str] = set()

    def add_candidate(path: Path) -> None:
        key = str(path)
        if key in seen:
            return
        seen.add(key)
        candidates.append(path)

    if raw_path.is_absolute():
        add_candidate(raw_path)
        if raw_path.suffix.lower() != ".json":
            add_candidate(raw_path.with_suffix(raw_path.suffix + ".json"))

    add_candidate(media_folder / raw_path.name)
    if raw_path.suffix.lower() != ".json":
        add_candidate(media_folder / f"{raw_path.name}.json")

    stem_name = raw_path.stem if raw_path.suffix.lower() == ".json" else raw_path.name
    add_candidate(media_folder / stem_name)
    if not str(stem_name).endswith(".json"):
        add_candidate(media_folder / f"{stem_name}.json")

    return candidates


def _suffix_matched_media_json_paths(project_id, image_name: str, settings=None) -> list[Path]:
    resolved_settings = settings or config
    paths = ProjectPaths(project_id=project_id, settings=resolved_settings)
    media_folder = paths.dir("media", prefer_existing=True)
    normalized_name = _normalize_media_reference(image_name)
    if not normalized_name or not media_folder.exists():
        return []

    raw_name = Path(normalized_name).name
    raw_name_cf = raw_name.casefold()
    suffix_names = {raw_name_cf}
    if not raw_name_cf.endswith(".json"):
        suffix_names.add(f"{raw_name_cf}.json")

    matches: list[Path] = []
    for candidate in sorted(media_folder.glob("*")):
        if not candidate.is_file():
            continue
        candidate_name_cf = candidate.name.casefold()
        if any(
            candidate_name_cf == suffix_name
            or candidate_name_cf.endswith(f"_{suffix_name}")
            for suffix_name in suffix_names
        ):
            matches.append(candidate)
    return matches


def _tokenize_media_match_text(text: str) -> set[str]:
    return {
        token
        for token in re.findall(r"[a-z0-9]+", str(text).lower())
        if len(token) >= 3
    }


def _candidate_media_context_text(candidate_path: Path) -> str:
    try:
        payload = json.loads(candidate_path.read_text(encoding="utf-8"))
    except Exception:
        return ""

    annotation = payload.get("image_annotation", {}) if isinstance(payload, dict) else {}
    parts = [
        annotation.get("short_description", ""),
        annotation.get("summary", ""),
        payload.get("image_key", ""),
        payload.get("filename", ""),
    ]
    return " ".join(str(part).strip() for part in parts if str(part).strip())


def _resolve_ambiguous_media_by_context(
    suffix_matches: list[Path],
    context_text: str | None,
) -> Path | None:
    if not context_text:
        return None

    context_tokens = _tokenize_media_match_text(context_text)
    if not context_tokens:
        return None

    scored_matches: list[tuple[int, Path]] = []
    for candidate_path in suffix_matches:
        candidate_tokens = _tokenize_media_match_text(_candidate_media_context_text(candidate_path))
        score = len(context_tokens & candidate_tokens)
        scored_matches.append((score, candidate_path))

    scored_matches.sort(key=lambda item: (-item[0], item[1].name))
    if not scored_matches or scored_matches[0][0] <= 0:
        return None
    if len(scored_matches) > 1 and scored_matches[0][0] == scored_matches[1][0]:
        return None
    return scored_matches[0][1]


def _build_slide_context_text(slide_payload: dict | None) -> str:
    if not isinstance(slide_payload, dict):
        return ""
    context_parts = []
    for key, value in slide_payload.items():
        if str(key) == TEMPLATE_KEY_FIELD:
            continue
        if "picture" in str(key).lower():
            continue
        if isinstance(value, str) and value.strip():
            context_parts.append(value.strip())
    return " ".join(context_parts)


def _resolve_media_json_path(project_id, image_name: str, settings=None, context_text: str | None = None) -> Path:
    candidate_paths = _candidate_media_json_paths(project_id, image_name, settings=settings)
    for candidate_path in candidate_paths:
        if candidate_path.is_file():
            return candidate_path

    suffix_matches = _suffix_matched_media_json_paths(project_id, image_name, settings=settings)
    if len(suffix_matches) == 1:
        return suffix_matches[0]
    if len(suffix_matches) > 1:
        contextual_match = _resolve_ambiguous_media_by_context(suffix_matches, context_text)
        if contextual_match is not None:
            return contextual_match
        raise ValueError(
            f"Ambiguous media reference '{image_name}'. Matches: "
            + ", ".join(str(path) for path in suffix_matches)
        )

    raise FileNotFoundError(
        "Media json file was not found. Checked: "
        + ", ".join(str(path) for path in candidate_paths)
    )


def get_image_path(project_id, image_name: str, settings=None, context_text: str | None = None) -> str:
    resolved_settings = settings or config
    paths = ProjectPaths(project_id=project_id, settings=resolved_settings)
    json_image_path = _resolve_media_json_path(
        project_id,
        image_name,
        settings=resolved_settings,
        context_text=context_text,
    )

    with open(json_image_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    image_data = data["image_base64"]
    header, base64_data = image_data.split(",", 1)
    ext = header.split(";")[0].split("/")[-1]
    img_bytes = base64.b64decode(base64_data)

    save_dir = paths.dir("images")
    save_dir.mkdir(parents=True, exist_ok=True)

    base_name = Path(json_image_path).stem
    if base_name.lower().endswith(f".{ext.lower()}"):
        save_file_name = base_name
    else:
        save_file_name = f"{base_name}.{ext}"
    save_path = save_dir / save_file_name

    with open(save_path, "wb") as img_file:
        img_file.write(img_bytes)

    return str(save_path)


def resolve_media_reference_name(project_id, image_name: str, settings=None, context_text: str | None = None) -> str:
    normalized_name = _normalize_media_reference(image_name)
    if not normalized_name or _is_sentinel_media(normalized_name):
        return normalized_name

    json_media_path = _resolve_media_json_path(
        project_id,
        normalized_name,
        settings=settings,
        context_text=context_text,
    )
    media_name = json_media_path.name
    if media_name.lower().endswith(".json"):
        media_name = media_name[:-5]
    return media_name


def fill_presentation(project_id, helper, template_dict, content_dict):
    content_dict = {
        i + 1: strip_slide_payload_metadata(v) if isinstance(v, dict) else v
        for i, v in enumerate(content_dict.values())
    }
    for slide_key, shapes in template_dict.items():
        slide_index = slide_key - 1
        print(f"Filling slide {slide_index + 1}")
        slide_payload = content_dict[slide_key]
        context_text = _build_slide_context_text(slide_payload)
        template_key = _template_key_from_shapes(shapes)
        is_story_with_image_points = _is_story_with_image_and_points_template(shapes)
        is_text_image_focus_repeat = template_key == _IMG_FOCUS_REPEAT_TEMPLATE_KEY
        handled_shape_keys = set()
        handled_shape_keys.update(
            _apply_adaptive_bullet_layout(
                helper=helper,
                slide_index=slide_index,
                shapes=shapes,
                slide_payload=slide_payload,
            )
        )
        handled_shape_keys.update(
            _apply_story_with_image_and_points_layout(
                helper=helper,
                slide_index=slide_index,
                shapes=shapes,
                slide_payload=slide_payload,
            )
        )
        handled_shape_keys.update(
            _apply_text_with_image_focus_repeat_layout(
                helper=helper,
                slide_index=slide_index,
                shapes=shapes,
                slide_payload=slide_payload,
            )
        )
        for shape_key, shape_info in shapes.items():
            if shape_key in handled_shape_keys:
                continue
            if str(shape_key).startswith("__"):
                continue
            if not isinstance(shape_info, dict) or "name" not in shape_info:
                continue
            shape_name = shape_info["name"]
            if shape_key not in slide_payload:
                continue
            new_text = slide_payload[shape_key]

            shape = helper.get_shape_by_name(slide_index, shape_name)
            if shape:
                if "picture" in shape_name.lower():
                    media_name = str(new_text).strip()
                    if not media_name or _is_sentinel_media(media_name):
                        remove_shape = getattr(helper, "remove_shape", None)
                        if callable(remove_shape):
                            remove_shape(shape)
                        continue

                    try:
                        img_path = get_image_path(
                            project_id,
                            media_name,
                            context_text=context_text,
                        )
                    except Exception as exc:
                        raise ValueError(
                            f"Failed media resolution on slide {slide_index + 1}, "
                            f"shape '{shape_name}', media '{media_name}'"
                        ) from exc

                    final_img_path = img_path
                    if str(shape_key) == "picture_placeholder" and is_story_with_image_points:
                        try:
                            final_img_path = _prepare_image_for_cover_box(
                                img_path,
                                target_width_emu=shape.width,
                                target_height_emu=shape.height,
                            )
                        except Exception as exc:
                            raise ValueError(
                                f"Failed image preprocessing on slide {slide_index + 1}, "
                                f"shape '{shape_name}', media '{media_name}'"
                            ) from exc
                    elif str(shape_key) == "picture_placeholder" and is_text_image_focus_repeat:
                        try:
                            final_img_path = _prepare_image_for_contain_box(
                                img_path,
                                target_width_emu=shape.width,
                                target_height_emu=shape.height,
                            )
                        except Exception as exc:
                            raise ValueError(
                                f"Failed image preprocessing on slide {slide_index + 1}, "
                                f"shape '{shape_name}', media '{media_name}'"
                            ) from exc

                    fill_image_placeholder(
                        helper,
                        slide_index=slide_index,
                        placeholder_name=shape_name,
                        image_path=final_img_path,
                    )
                else:
                    helper.replace_text_preserve_format(shape, new_text)


def keep_slides_by_index(helper, response, template, map_dict):
    canonical_slides = canonicalize_slide_payload(
        response,
        available_templates=template,
    )
    selected_slides = [
        str(slide_payload[TEMPLATE_KEY_FIELD]).strip()
        for slide_payload in canonical_slides.values()
    ]
    slide_numbers = [map_dict[name] for name in selected_slides]

    slides_indices_keep = sorted(slide_numbers)

    if not slides_indices_keep:
        raise ValueError("No valid slide_X keys found in new_r.")

    total_slides = len(helper.prs.slides)
    content_count = total_slides - 1

    invalid = [i for i in slides_indices_keep if i < 1 or i > content_count]
    if invalid:
        raise ValueError(
            f"Invalid indices {invalid}. Valid range is 1..{content_count} (last slide excluded)."
        )

    zero_idx = [i - 1 for i in slides_indices_keep]
    keep_set = set(zero_idx)

    for i in reversed(range(content_count)):
        if i not in keep_set:
            helper.delete_slide(i)

    updated_template = {
        i: {
            _TEMPLATE_KEY_META_FIELD: key,
            **template[key]["placeholders"],
        }
        for i, key in enumerate(selected_slides, 1)
    }

    return updated_template, slide_numbers


def compute_reorder_list(agent_selected):
    sorted_agent_selected = sorted(agent_selected)
    original_to_new = {
        original_idx: new_idx + 1
        for new_idx, original_idx in enumerate(sorted_agent_selected)
    }

    reorder = [original_to_new[i] for i in agent_selected]

    return reorder


def reorder_slides(helper, new_order):
    total_slides = len(helper.prs.slides)
    content_count = total_slides - 1

    zero_idx = [i - 1 for i in new_order]

    if len(zero_idx) != content_count:
        raise ValueError(
            f"new_order must contain {content_count} indices (excluding last slide)"
        )

    if sorted(zero_idx) != list(range(content_count)):
        raise ValueError("new_order must be a permutation of content slide indices")

    sldIdLst = helper.prs.slides._sldIdLst
    sldId_elems = list(sldIdLst)

    content_slides = sldId_elems[:-1]
    last_slide = sldId_elems[-1]

    sldIdLst.clear()
    for idx in zero_idx:
        sldIdLst.append(content_slides[idx])
    sldIdLst.append(last_slide)
